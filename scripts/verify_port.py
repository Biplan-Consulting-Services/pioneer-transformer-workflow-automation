# -*- coding: utf-8 -*-
"""Compare FRM10-12's TableOrders against Order Items, value by value.

    python scripts/verify_port.py            # summary
    python scripts/verify_port.py --detail   # plus up to 5 example rows per column

`rediff_units.py` answers "is every unit there". This answers the harder question:
"does every unit hold the same DATA it held in the workbook".

WHERE THE MAPPING COMES FROM
  Not from a list written by hand here -- from the live flow definition itself, the
  newest pulled version. For each `item/<target>` on `UpdateOrderItem` it finds the
  Excel column the expression reads, so this cannot drift away from what the flow
  actually does. Of 130 mappings:

      49  read an Excel column   <- the only ones this can check
      63  read a SharePoint parent through a lookup or variable
      18  are computed (Title, MappedLocation, ItemStatus, PoItemNumber, ...)

  The 81 non-Excel ones are out of scope by definition: their source is not the
  workbook, so the workbook cannot confirm them. `x7_verify_parent_sync.js` is the
  tool for those.

THE TRANSFORMS ARE APPLIED, NOT IGNORED
  A raw string compare would report thousands of false differences, because the flow
  rewrites values on the way in. Each is modelled from the expression actually found:

      @not(equals(<col>, ''))     marker -> boolean. `R`/`x`/`Y` in the workbook is
                                  TRUE on the list, and BLANK is FALSE. This is the
                                  single biggest group.
      toLower(<col>)              case-insensitive compare
      plain @item()?['<col>']     compare as text

  Numbers and dates are normalised before comparison: the export writes `4,260` and
  `2,000` with thousands separators and renders dates in its own format, neither of
  which is a difference in the data.

⚠️ READ THE ZERO CAREFULLY. A column can only be checked where BOTH sides have a
   value to compare. Columns that are empty on both sides agree trivially and are
   reported separately, so "0 mismatches" on a column nobody fills means nothing.
"""
import argparse, csv, io, glob, json, os, re, sys, warnings

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
csv.field_size_limit(min(2 ** 31 - 1, sys.maxsize))

UNIT = re.compile(r"^[A-Za-z0-9_]+-\d+/\d+$")


def newest(pat):
    f = sorted(glob.glob(os.path.join(ROOT, pat)))
    if not f:
        sys.exit("nothing matching %s" % pat)
    return f[-1]


# ---------------------------------------------------------------- the mapping
def mapping():
    """{Order Items target -> (excel column, kind)} straight out of the live flow."""
    from flow_version import read_any, definition
    p = sorted(glob.glob(os.path.join(
        ROOT, "workflow-data", "Order Items - excel transfer flow", "v0*.json")))[-1]
    d = definition(read_any(p)[0])

    found = {}

    def walk(o):
        if isinstance(o, dict):
            for k, v in o.items():
                if k == "UpdateOrderItem":
                    par = (v.get("inputs") or {}).get("parameters") or {}
                    for t, e in par.items():
                        if t.startswith("item/") and isinstance(e, str):
                            found[t] = e
                walk(v)
        elif isinstance(o, list):
            for x in o:
                walk(x)
    walk(d)

    out = {}
    for tgt, e in found.items():
        srcs = set(re.findall(r"item\(\)\?\['([^']+)'\]", e))
        if len(srcs) != 1:
            continue                      # parent-sourced or computed: not ours
        if re.search(r"body\('|variables\('|outputs\('", e):
            continue
        col = srcs.pop()
        if re.fullmatch(r"@item\(\)\?\['[^']+'\]", e.strip()):
            kind = "text"
        elif "not(equals(" in e and "'')" in e:
            kind = "bool"
        elif "'Completed'" in e or "'Pending'" in e:
            # `if(empty(<a date column>), 'Pending', 'Completed')`. The workbook holds
            # the stage's DATE; the list holds a status derived from whether that date
            # exists. Comparing them as text reports every single row as different,
            # which is what the first run of this script did across six columns.
            kind = "date2status"
        elif "toLower(" in e:
            kind = "lower"
        else:
            kind = "expr"
        out[tgt] = (col, kind)
    return out, os.path.basename(p)


# ---------------------------------------------------------------- the two sides
def workbook():
    import openpyxl
    from openpyxl.utils import range_boundaries
    p = newest("workbooks/FRM10-12 2*.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb["Orders"]
    c1, r1, c2, r2 = range_boundaries(ws.tables["TableOrders"].ref)
    hdr = [str(ws.cell(row=r1, column=c).value or "").strip() for c in range(c1, c2 + 1)]
    key = None
    for i, c in enumerate(range(c1, c2 + 1)):
        n = sum(1 for rr in range(r1 + 1, r2 + 1)
                if UNIT.match(str(ws.cell(row=rr, column=c).value or "").strip()))
        if n > 500:
            key = i
            break
    rows = {}
    for rr in range(r1 + 1, r2 + 1):
        vals = [ws.cell(row=rr, column=c).value for c in range(c1, c2 + 1)]
        u = str(vals[key] or "").strip()
        if u:
            rows[u] = dict(zip(hdr, vals))
    wb.close()
    return os.path.basename(p), rows


def order_items():
    p = newest("sharepoint-lists/Order Items 2*.csv")
    r = csv.reader(io.open(p, encoding="utf-8-sig", newline=""))
    first = next(r)
    hdr = next(r) if first and first[0].startswith("ListSchema=") else first
    rows = {}
    key = hdr.index("Unit ID")
    for x in r:
        u = (x[key] or "").strip()
        if u:
            rows[u] = dict(zip(hdr, x))
    return os.path.basename(p), hdr, rows


def slug(s):
    """Internal name to comparable form: decode _x00XX_, drop everything else."""
    s = re.sub(r"_x00([0-9a-fA-F]{2})_", lambda m: chr(int(m.group(1), 16)), s)
    return re.sub(r"[^a-z0-9]", "", s.lower())


# ---------------------------------------------------------------- comparison
TRUE = {"true", "yes", "oui", "1"}


RICHTEXT = re.compile(r'^<div class="ExternalClass[0-9A-F]+">(.*)</div>$', re.S)


def norm(v):
    if v is None:
        return ""
    # datetime/date/time before str(): Excel hands back objects, the export strings
    if hasattr(v, "year"):
        return "%04d-%02d-%02d" % (v.year, v.month, v.day)
    if hasattr(v, "hour") and not hasattr(v, "year"):
        return "0" if (v.hour, v.minute, v.second) == (0, 0, 0) else str(v)
    s = str(v).strip()
    # SharePoint wraps multi-line text in its own div on the way into Power Query;
    # the list holds the inner text. Same content, different envelope.
    m = RICHTEXT.match(s)
    if m:
        s = m.group(1).strip()
    if s.lower() in ("none", "nan"):
        return ""
    # 4,260 / 2,000.0 -> 4260 ; dates -> ISO date part
    t = s.replace(",", "")
    try:
        f = float(t)
        return str(int(f)) if f == int(f) else str(f)
    except ValueError:
        pass
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return "%s-%s-%s" % m.groups()
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--detail", action="store_true")
    a = ap.parse_args()

    maps, defn = mapping()
    wbname, wbrows = workbook()
    oiname, oihdr, oirows = order_items()
    by_slug = {slug(h): h for h in oihdr}

    print("mapping from : %s" % defn)
    print("workbook     : %-44s %5d units" % (wbname, len(wbrows)))
    print("Order Items  : %-44s %5d units" % (oiname, len(oirows)))
    both = [u for u in wbrows if u in oirows]
    print("compared on  : %d units present in both\n" % len(both))

    unresolved, results = [], []
    for tgt, (col, kind) in sorted(maps.items()):
        name = tgt[5:].split("/")[0]
        oh = by_slug.get(slug(name))
        if oh is None or col not in next(iter(wbrows.values())):
            unresolved.append((tgt, col, "no list column" if oh is None else "no workbook column"))
            continue
        same = diff = bothblank = onlyone = 0
        ex = []
        for u in both:
            w, o = wbrows[u].get(col), oirows[u].get(oh)
            if kind == "date2status":
                raw = str(w).strip().lower() if w is not None else ""
                # `ec` is en cours in the workbook, and the flow translates it rather
                # than treating it as a date. It is the only non-date value the stage
                # columns carry, and FRM11's purge rule depends on the same code.
                wv = "in progress" if raw == "ec" else ("completed" if norm(w) else "")
                ov = str(o).strip().lower()
                if wv == "" and ov in ("", "pending"):
                    bothblank += 1
                    continue
            elif kind == "bool":
                wv = "true" if norm(w) else "false"
                ov = "true" if str(o).strip().lower() in TRUE else "false"
                if not norm(w) and str(o).strip() == "":
                    bothblank += 1
                    continue
            else:
                wv, ov = norm(w), norm(o)
                if kind == "lower":
                    wv, ov = wv.lower(), ov.lower()
                if wv == "" and ov == "":
                    bothblank += 1
                    continue
                if wv == "" or ov == "":
                    onlyone += 1
                    if len(ex) < 5:
                        ex.append((u, w, o))
                    continue
            if wv == ov:
                same += 1
            else:
                diff += 1
                if len(ex) < 5:
                    ex.append((u, w, o))
        results.append((name, col, kind, same, diff, onlyone, bothblank, ex))

    print("%-34s %-6s %7s %6s %8s %7s" %
          ("Order Items column", "kind", "match", "DIFF", "one-side", "blank"))
    bad = 0
    for name, col, kind, same, diff, one, bb, ex in results:
        flag = "  <<<" if diff else ""
        if diff:
            bad += 1
        print("%-34s %-6s %7d %6d %8d %7d%s" % (name[:34], kind, same, diff, one, bb, flag))
        if a.detail and ex:
            for u, w, o in ex:
                print("        %-14s workbook=%-22r list=%r" % (u, w, o))

    print("\n%d column(s) compared, %d with at least one real difference" %
          (len(results), bad))
    tot_d = sum(r[4] for r in results)
    tot_s = sum(r[3] for r in results)
    print("%d values agree, %d differ" % (tot_s, tot_d))
    if unresolved:
        print("\nnot compared (column not present on one side):")
        for t, c, why in unresolved:
            print("   %-34s %-28s %s" % (t, c, why))
    return 1 if tot_d else 0


if __name__ == "__main__":
    sys.exit(main())
