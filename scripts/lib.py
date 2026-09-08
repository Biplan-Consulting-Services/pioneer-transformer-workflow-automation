# -*- coding: utf-8 -*-
"""Shared loaders + normalisation for the workbook-vs-list comparison.

Every normalisation rule here exists because this repo has already been burned
by it once; the comment on each says which."""
import re, datetime, json, io, os, sys
sys.path.insert(0, os.path.abspath("scripts"))
from load_exports import load
import openpyxl

EPOCH = datetime.date(1899,12,30)   # the flow uses addDays('1899-12-30', int(x))

def xl_table(path, table_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        for ws in wb.worksheets:
            for n, ref in (ws.tables or {}).items():
                if n.lower() == table_name.lower():
                    cells = ws[ref]
                    hdr = [str(c.value).strip() if c.value is not None else "" for c in cells[0]]
                    rows = []
                    for r in cells[1:]:
                        d = {hdr[i]: r[i].value for i in range(len(hdr))}
                        rows.append(d)
                    return hdr, rows
        raise SystemExit("table %s not found in %s" % (table_name, path))
    finally:
        wb.close()

JSONREF = re.compile(r'^\s*\[\s*\{.*"Value"\s*:\s*"([^"]*)".*\}\s*\]\s*$', re.S)

def norm(v, kind="text"):
    """Return a comparable scalar, or None for 'no value'."""
    if v is None: return None
    if isinstance(v, bool):                       # openpyxl gives real bools
        return "true" if v else "false"
    if isinstance(v, (datetime.datetime, datetime.date)):
        d = v.date() if isinstance(v, datetime.datetime) else v
        return d.isoformat()
    s = str(v).strip()
    if s == "" : return None
    # R19 trap: a lookup/multichoice exports (and, per R22, is even WRITTEN) as
    # [{"@odata.type":...,"Value":"X"}] -- unwrap to X.
    m = JSONREF.match(s)
    if m: s = m.group(1).strip()
    if s == "": return None
    low = s.lower()
    if low in ("yes","true","1","oui"):  return "true"
    if low in ("no","false","0","non"):  return "false"
    if kind == "date":
        # Excel serial
        if re.fullmatch(r'\d+(\.\d+)?', s):
            try: return (EPOCH + datetime.timedelta(days=float(s))).date().isoformat()
            except Exception: pass
        for f in ("%Y-%m-%d","%m/%d/%Y","%d/%m/%Y","%Y-%m-%dT%H:%M:%SZ","%m/%d/%Y %H:%M","%Y-%m-%d %H:%M:%S"):
            try: return datetime.datetime.strptime(s.split("T")[0] if f=="%Y-%m-%d" else s, f).date().isoformat()
            except Exception: pass
        m2 = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
        if m2: return "%s-%s-%s" % m2.groups()
        return low
    if kind == "num":
        # R19 trap: French vs English decimal separator, plus stray spaces
        t = s.replace(" ", "").replace("\u00a0","")
        if re.fullmatch(r'-?\d+(?:[.,]\d+)?', t):
            t = t.replace(",", ".")
            try:
                f = float(t)
                return str(int(f)) if f == int(f) else ("%.6f" % f).rstrip("0").rstrip(".")
            except Exception: pass
        return low
    return re.sub(r'\s+', " ", low)

MARKERS = {"ec","tbd"}      # guard values the flow deliberately turns into null
def is_marker(v):
    return v is not None and str(v).strip().lower() in MARKERS
