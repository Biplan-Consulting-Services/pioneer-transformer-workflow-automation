# -*- coding: utf-8 -*-
"""Every unit the archive calls DONE must be GONE from the `Order Items` list.

    python scripts/verify_archive_done_not_in_list.py
    python scripts/verify_archive_done_not_in_list.py --csv reports/archive-done-still-live.csv

DONE is defined on the ARCHIVE side only:  `Location = LI`  AND  `Delivery Date` holds
a real date.  The list's own `Location` / `Item Status` are NOT part of the test -- a
delivered unit is delivered whatever the list still claims, and the whole point is to
catch rows whose list state never caught up.  The list side answers one question:
is this unit id present.

WHY THIS IS THE INVERSE OF `rediff_units.py`
  That script's note says a list-only row "is not an error -- units leave the workbook
  when they are archived, and stay on the list", and counted 126 of them on 2026-09-09,
  "all confirmed Location = LI with a real delivery date".  That was true while the list
  was the live system and nothing pruned it.  This script is the cleanup check that turns
  the same population into a worklist: a unit that is finished in the archive has no
  business still sitting in the active list.

  So a NON-ZERO result here is not necessarily a bug in the flow.  It is a set of rows to
  decide about -- retire, or explain.

⚠️ READ THE ZERO CAREFULLY -- TWO WAYS IT CAN LIE
  1. DONE is a conjunction.  An archive row parked at `LI` with NO delivery date is not
     matched by it and will not be reported, however finished it looks.  That population
     is printed every run, so the zero cannot quietly mean "nothing qualified".

  2. 🔴 AN EXPORT FOLLOWS THE SELECTED VIEW, NOT THE LIST.  A view that filters rows out
     makes them invisible here, and an invisible row reads as "already gone" -- a FALSE
     PASS, on exactly the population being hunted.  The export of 2026-09-16 10:01 is
     this case: 82 columns and `Item Status = Active` on all 1079 rows, where the
     2026-09-11 export had 142 columns and 102 `Delivered` rows.  Every already-retired
     unit was filtered out of the file.

     The file proves this about itself.  Its `ListSchema=` record declares every field on
     the LIST (152), while the header carries only what the VIEW selected (82).  That gap
     is checked on every run and a wide one is called out.  `stamp_exports.py` catches the
     same thing by row count; this catches it by shape, before the answer is believed.

     Export from **All Items** for a conclusive run.

COLUMNS ARE FOUND BY VALUE, NOT BY LABEL
  The unit id lives under `Order` in the archive and `Unit ID` on the list, and both
  names have moved before.  x8 matched `ServerRedirectedEmbedUrl` by name and printed 24
  blank paths; the lesson stuck.  Both sides score every column by how many of its values
  actually look like a unit id -- `21670-5/5`, or `21611-1/1 SA` -- and take the winner.
  `Location` and `Delivery Date` ARE looked up by header, because nothing else in the
  archive distinguishes them, but each is then sanity-checked against its own values and
  the run aborts rather than testing the wrong column.

  ⚠️ That shape test picks the COLUMN. It does not filter the ROWS. An earlier cut of
  this script required every id to match the pattern and quietly dropped 360 archive
  rows, 7 of them DONE -- `19394-1/2SA` (no space), `19515-W4-14`, `E-21000-1/1`. Real
  units with untidy ids are exactly the ones likeliest to be missed by a cleanup, so
  every non-empty id is carried, whatever it looks like.

THE ` SA` SUFFIX IS PART OF THE ID
  Both sides carry SA jobs as `21444-1/3 SA`, so ids compare EXACTLY and no stripping is
  wanted: `21444-1/3` and `21444-1/3 SA` are different units and must not collide.
  Because the archive also spells some of them `19394-1/2SA`, a second pass reports ids
  that match only once case and whitespace are ignored -- not as a verdict, as a warning
  that a spelling drift could be hiding a row. It finds 0 today.
"""
import argparse
import collections
import csv
import datetime
import io
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
csv.field_size_limit(min(2 ** 31 - 1, sys.maxsize))

import openpyxl  # noqa: E402

try:                                    # a cp1252 console cannot print ✅/⚠️
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DEFAULT_ARCHIVE = os.path.join(ROOT, "workbooks", "Archive active.xlsx")
DEFAULT_LIST = os.path.join(ROOT, "sharepoint-lists", "Order Items.csv")
# The unit-grain sheet. `Archive BO` / `FRM11` / `FRM13` are other grains entirely and
# would silently compare the wrong population, so this is named, not guessed.
DEFAULT_SHEET = "Archive FRM10-12"

UNIT = re.compile(r"^[A-Za-z0-9_]+-\d+/\d+(?: SA)?$")
DONE_LOCATION = "LI"
EPOCH = datetime.date(1899, 12, 30)   # the flow's own addDays('1899-12-30', int(x))


def stamp(path):
    t = datetime.datetime.fromtimestamp(os.path.getmtime(path))
    return "%s   (saved %s)" % (os.path.relpath(path, ROOT), t.strftime("%Y-%m-%d %H:%M"))


def pick_unit_column(header, rows, where):
    """The column whose VALUES look like unit ids, not the one whose name does."""
    best, hits = None, -1
    for i in range(len(header)):
        n = sum(1 for r in rows if i < len(r) and UNIT.match(str(r[i] or "").strip()))
        if n > hits:
            best, hits = i, n
    if hits <= 0:
        sys.exit("%s: no column holds anything shaped like a unit id" % where)
    return best, hits


def as_date(v):
    """A real date, or None. Excel serials and exported strings both count."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        try:
            return EPOCH + datetime.timedelta(days=float(v))
        except Exception:
            return None
    s = str(v).strip()
    if not s:
        return None
    if re.fullmatch(r"\d+(\.\d+)?", s):
        try:
            return EPOCH + datetime.timedelta(days=float(s))
        except Exception:
            return None
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%SZ",
              "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(s, f).date()
        except Exception:
            pass
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return datetime.date(*(int(x) for x in m.groups()))
    return None


def read_archive(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            sys.exit("archive has no sheet %r; it has %s" % (sheet, wb.sheetnames))
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if len(rows) < 2:
        sys.exit("archive sheet %r is empty" % sheet)

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    body = rows[1:]

    u, hits = pick_unit_column(header, body, "archive")

    def column(name, test, what):
        if name not in header:
            sys.exit("archive sheet %r has no %r column; header is %s" % (sheet, name, header))
        i = header.index(name)
        ok = sum(1 for r in body if i < len(r) and test(r[i]))
        if ok == 0:
            sys.exit("archive column %r holds no %s -- wrong column, refusing to judge "
                     "DONE from it" % (name, what))
        return i, ok

    # Looked up by name, then made to prove itself. A Location column with no location
    # codes in it, or a date column with no dates, means the sheet moved under us.
    loc, _ = column("Location", lambda v: str(v or "").strip().upper() == DONE_LOCATION,
                    "%r values" % DONE_LOCATION)
    dd, _ = column("Delivery Date", lambda v: as_date(v) is not None, "dates")

    done, li_no_date, odd, seen = {}, [], [], collections.Counter()
    for r in body:
        uid = str(r[u]).strip() if u < len(r) and r[u] is not None else ""
        if not uid:
            continue
        seen[uid] += 1
        is_li = str(r[loc] or "").strip().upper() == DONE_LOCATION if loc < len(r) else False
        date = as_date(r[dd]) if dd < len(r) else None
        if is_li and date is not None:
            done[uid] = date
            if not UNIT.match(uid):
                odd.append(uid)          # carried anyway -- see the docstring
        elif is_li:
            li_no_date.append(uid)       # LI but undated -> NOT done by the rule

    return {
        "unit_column": header[u], "unit_hits": hits,
        "rows": len(body), "done": done, "li_no_date": li_no_date,
        "odd": odd, "dupes": {k: v for k, v in seen.items() if v > 1},
    }


def read_list(path):
    """SharePoint 'Export to CSV': a giant ListSchema=... record, then the real header."""
    raw = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            raw = io.open(path, encoding=enc, newline="").read()
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        sys.exit("cannot decode %s" % path)

    rows = list(csv.reader(io.StringIO(raw)))
    if not rows:
        sys.exit("%s is empty" % path)

    # How many fields the LIST has, per the export's own schema record. The record spans
    # many lines AND many csv fields, so join it back together before counting.
    declared = None
    if rows and rows[0] and str(rows[0][0]).startswith("ListSchema="):
        declared = ",".join(str(c) for c in rows[0]).count("<Field ")
        rows = rows[1:]                  # the first RECORD, which spans many lines
    header, body = rows[0], rows[1:]
    body = [r for r in body if any(str(c).strip() for c in r)]

    u, hits = pick_unit_column(header, body, "list")

    ids, blank, odd = {}, 0, []
    for r in body:
        uid = str(r[u]).strip() if u < len(r) and r[u] is not None else ""
        if not uid:
            blank += 1
            continue
        if not UNIT.match(uid):
            odd.append(uid)              # carried anyway -- see the docstring
        ids.setdefault(uid, r)

    # A column with exactly one distinct value across a thousand rows is a filter, not
    # a coincidence -- and `Item Status` is the one the retiring view filters on.
    single = {}
    for c in ("Item Status", "Location"):
        if c in header:
            i = header.index(c)
            vals = set(str(r[i]).strip() for r in body if i < len(r) and str(r[i]).strip())
            if len(vals) == 1:
                single[c] = vals.pop()

    return {"unit_column": header[u], "unit_hits": hits, "header": header,
            "rows": len(body), "ids": ids, "blank": blank, "odd": odd,
            "declared": declared, "exported": len(header), "single": single}


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--archive", default=DEFAULT_ARCHIVE)
    ap.add_argument("--list", dest="lst", default=DEFAULT_LIST)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--csv", help="also write the offenders to this file")
    a = ap.parse_args()

    for p in (a.archive, a.lst):
        if not os.path.exists(p):
            sys.exit("missing input: %s" % p)

    print("archive : %s   [%s]" % (stamp(a.archive), a.sheet))
    print("list    : %s" % stamp(a.lst))
    print("")

    arc = read_archive(a.archive, a.sheet)
    lst = read_list(a.lst)

    print("archive : %d rows, unit id read from %r (%d unit-shaped values)"
          % (arc["rows"], arc["unit_column"], arc["unit_hits"]))
    print("          %d DONE  (Location = %s AND Delivery Date set)"
          % (len(arc["done"]), DONE_LOCATION))
    print("          %d at %s with NO delivery date -- NOT tested, see the docstring"
          % (len(arc["li_no_date"]), DONE_LOCATION))
    if arc["odd"]:
        print("          %d DONE id(s) are untidy but still tested: %s"
              % (len(arc["odd"]), ", ".join(sorted(set(arc["odd"]))[:6])))
    if arc["dupes"]:
        print("          ⚠️  %d unit id(s) appear more than once: %s"
              % (len(arc["dupes"]), ", ".join(sorted(arc["dupes"])[:5])))
    print("list    : %d rows, unit id read from %r (%d unit-shaped values)"
          % (lst["rows"], lst["unit_column"], lst["unit_hits"]))
    if lst["blank"]:
        print("          ⚠️  %d row(s) carry no unit id at all and cannot be matched"
              % lst["blank"])
    if lst["odd"]:
        print("          %d untidy id(s), still tested: %s"
              % (len(lst["odd"]), ", ".join(sorted(set(lst["odd"]))[:6])))
    print("")

    # ---- is this export the whole list, or one view of it? ----------------------
    view_shaped = []
    if lst["declared"] and lst["exported"] < lst["declared"] * 0.75:
        view_shaped.append("it carries %d columns where its own ListSchema declares %d "
                           "fields on the list" % (lst["exported"], lst["declared"]))
    for c, v in sorted(lst["single"].items()):
        view_shaped.append("every one of its %d rows has %s = %r" % (lst["rows"], c, v))

    if view_shaped:
        print("🔴 THIS EXPORT LOOKS LIKE A VIEW, NOT THE WHOLE LIST")
        for r in view_shaped:
            print("     - %s" % r)
        print("   Rows the view filters out are invisible here, and an invisible row"
              "\n   reads as 'already gone'. A PASS below is NOT conclusive; the units"
              "\n   most likely to be missing are the ones already retired -- exactly the"
              "\n   population this check is about. Re-export from All Items to settle it.")
        print("")

    offenders = sorted(set(arc["done"]) & set(lst["ids"]))

    # Near misses: same unit, different spelling. Not a verdict -- a warning that an
    # offender could be hiding behind `19394-1/2SA` vs `19394-1/2 SA`.
    def loose(s):
        return re.sub(r"\s+", "", s).upper()
    lmap = {}
    for k in lst["ids"]:
        lmap.setdefault(loose(k), k)
    near = sorted((a, lmap[loose(a)]) for a in arc["done"]
                  if a not in lst["ids"] and loose(a) in lmap)
    if near:
        print("⚠️  %d DONE unit(s) match a list row only once case/spacing is ignored."
              "\n    Not counted below -- confirm whether these are the same unit:" % len(near))
        for a, b in near:
            print("      archive %-20s  list %s" % (a, b))
        print("")

    # The list's own columns are CONTEXT for triage, never part of the verdict.
    hdr = lst["header"]
    ctx = [c for c in ("Location", "Step Status", "Item Status", "Delivery Date") if c in hdr]

    if not offenders:
        print("%s none of the %d units the archive calls DONE appears in this export."
              % ("⚠️  INCONCLUSIVE —" if view_shaped else "✅ PASS —", len(arc["done"])))
        if view_shaped:
            print("   Read that as 'not in this view', not as 'not on the list'.")
    else:
        print("❌ %d unit(s) are DONE in the archive and STILL on the list%s:\n"
              % (len(offenders), " (at least — see the view warning above)" if view_shaped else ""))
        w = max(len(u) for u in offenders)
        for uid in offenders:
            row = lst["ids"][uid]
            bits = ["%-*s  archive Delivery Date %s" % (w, uid, arc["done"][uid].isoformat())]
            for c in ctx:
                i = hdr.index(c)
                v = str(row[i]).strip() if i < len(row) and row[i] is not None else ""
                bits.append("list %s=%s" % (c, v if v else "(blank)"))
            print("  " + "  |  ".join(bits))
        print("\n  (the `list ...` values are context for triage only -- DONE is decided"
              "\n   entirely by the archive, so a row still marked Active here is exactly"
              "\n   the case this check exists to surface)")

    if a.csv:
        out = a.csv if os.path.isabs(a.csv) else os.path.join(ROOT, a.csv)
        d = os.path.dirname(out)
        if d and not os.path.isdir(d):
            os.makedirs(d)
        with io.open(out, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["Unit ID", "Archive Delivery Date"] + ["List " + c for c in ctx])
            for uid in offenders:
                row = lst["ids"][uid]
                w.writerow([uid, arc["done"][uid].isoformat()] +
                           [(str(row[hdr.index(c)]).strip()
                             if hdr.index(c) < len(row) and row[hdr.index(c)] is not None else "")
                            for c in ctx])
        rel = os.path.relpath(out, ROOT)
        print("\nwrote %s  (%d row(s))" % (out if rel.startswith("..") else rel, len(offenders)))

    return 1 if offenders else 0


if __name__ == "__main__":
    sys.exit(main())
