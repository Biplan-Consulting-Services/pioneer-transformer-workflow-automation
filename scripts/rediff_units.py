# -*- coding: utf-8 -*-
"""Runbook 2.3 — re-diff units both directions, workbook vs `Order Items`.

    python scripts/rediff_units.py

Reads the newest `workbooks/FRM10-12 *.xlsx` and the newest
`sharepoint-lists/Order Items *.csv`, and reports units present in one and not the other.

**After the final run, both directions should be 0** (list-only rows excepted, see below).
A non-zero workbook-only count no longer has a known explanation: every previously known
cause was resolved on 2026-09-08/09 and the three expected creations are named below.

WHY THE UNIT COLUMN IS DISCOVERED, NOT NAMED
  On the list it is `Title`, displayed as `Unit ID`, and the export carries that header
  **twice** (Title plus its linked-to-item twin, identical on every row). In the workbook
  it is a `TableOrders` column whose name has moved before. So both sides pick the column
  whose values actually look like a unit id, `21814-1/11`, rather than trusting a header.
  That is the same lesson as x8 matching `ServerRedirectedEmbedUrl` by name and printing
  24 blank paths: score by value, not by label.

WHAT A LIST-ONLY ROW MEANS
  Not an error. Units leave the workbook when they are archived, and stay on the list.
  The 2026-09-09 audit had 126 of them, all confirmed `Location = LI` with a real delivery
  date. The number only grows. It is workbook-only that must reach zero.
"""
import csv, glob, io, os, re, sys, warnings

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
csv.field_size_limit(min(2 ** 31 - 1, sys.maxsize))

UNIT = re.compile(r"^[A-Za-z0-9_]+-\d+/\d+$")

# Created by the final run because their Order row exists; see runbook 2.3.
EXPECTED_CREATES = {"20877R1-1/1", "P20002-1/1", "P21911_A-1/1"}


def newest(pattern):
    f = sorted(glob.glob(os.path.join(ROOT, pattern)))
    if not f:
        sys.exit("no file matching %s" % pattern)
    return f[-1]


def from_list():
    p = newest("sharepoint-lists/Order Items 2*.csv")
    r = csv.reader(io.open(p, encoding="utf-8-sig", newline=""))
    first = next(r)
    hdr = next(r) if first and first[0].startswith("ListSchema=") else first
    rows = list(r)
    best, hits = None, -1
    for i in range(len(hdr)):
        n = sum(1 for x in rows if i < len(x) and UNIT.match((x[i] or "").strip()))
        if n > hits:
            best, hits = i, n
    return os.path.basename(p), set(
        (x[best] or "").strip() for x in rows if best < len(x) and (x[best] or "").strip())


def from_workbook():
    import openpyxl
    from openpyxl.utils import range_boundaries
    p = newest("workbooks/FRM10-12 2*.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb["Orders"]
    c1, r1, c2, r2 = range_boundaries(ws.tables["TableOrders"].ref)
    best, hits = None, -1
    for c in range(c1, c2 + 1):
        n = sum(1 for rr in range(r1 + 1, r2 + 1)
                if UNIT.match(str(ws.cell(row=rr, column=c).value or "").strip()))
        if n > hits:
            best, hits = c, n
    out = set()
    for rr in range(r1 + 1, r2 + 1):
        v = str(ws.cell(row=rr, column=best).value or "").strip()
        if v:
            out.add(v)
    wb.close()
    return os.path.basename(p), out


def main():
    lname, lunits = from_list()
    wname, wunits = from_workbook()
    wonly = sorted(wunits - lunits)
    lonly = sorted(lunits - wunits)

    print("list      %-44s %5d units" % (lname, len(lunits)))
    print("workbook  %-44s %5d units" % (wname, len(wunits)))
    print("in both                                                 %5d" % len(wunits & lunits))
    print()
    print("workbook only (MUST be 0 after the run)                 %5d" % len(wonly))
    for u in wonly:
        tag = "  <- expected, the run creates it" if u in EXPECTED_CREATES else \
              "  <- \U0001F534 NOT a known case"
        print("      %-16s%s" % (u, tag))
    print()
    print("list only (archived out of the workbook, normal)        %5d" % len(lonly))
    if lonly:
        print("      %s%s" % (", ".join(lonly[:8]), " ..." if len(lonly) > 8 else ""))

    unknown = [u for u in wonly if u not in EXPECTED_CREATES]
    if unknown:
        print("\n\U0001F534 %d workbook unit(s) with no known explanation. Runbook 2.3: "
              "anything here is a NEW problem." % len(unknown))
        return 1
    if wonly:
        print("\nAll %d workbook-only units are the expected creations. Re-run this "
              "after the run; it should then be 0." % len(wonly))
    else:
        print("\n✓ Zero workbook-only units. 2.3 passes.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
