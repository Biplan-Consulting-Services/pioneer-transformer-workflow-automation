#!/usr/bin/env python3
"""Validate a rebuilt Archive active copy against the baseline copy and the live SharePoint mirror.

    python scripts/check_archive_rebuild.py --base <baseline.xlsx> --new <rebuilt.xlsx> [--report out.md]
        [--verbose] [--strict] [--first-build] [--today YYYY-MM-DD]

Plan: ~/.claude/plans/rerunning-the-dry-run-snug-lighthouse.md, section D (decisions D2/D3/D4/D6).
Column spec: docs/archive-frm10-12-column-recipe.md.

Live SharePoint state comes ONLY from the mirror tool's own output, sharepoint-lists/mirror/live/*.csv
(refresh it first: scripts/Refresh-SharePointMirror.ps1). Never the manual exports.

Checks (each PASS / WARN / SKIP / FAIL; exit 1 on any FAIL):
  1 Shape            FRM11/FRM13 headers + cells untouched; FRM10-12 = the pinned 93 then the 18 BO
                     detail columns (D2); FRM10-12 keeps table GUID {17014A43-...}; no duplicate Order.
  2 FRM10-12 cells   every cell of every Order in both copies: identical / cleaner / newer (= the mirror's
                     value through the recipe) / intended (D3 columns, E7 orders' dates) / REGRESSION.
  3 Date columns     no text in a date column (name contains "date", minus the refresher's free-text
                     list; list tables: the catalog's date-only columns). "EC" allowed in the six stage
                     dates; text identical to the baseline's same cell is legacy (WARN, not FAIL).
  4 List tables      the five new tables vs the mirror: Id sets, date-only = first 10 chars, lookup
                     display values, every mirror column present.
  5 Additivity       per archive table: rows never drop, every baseline key still there.
  6 Readers          columns FRM11 / FRM13 / viewer purge / Query2 / PriceReg read exist; Location codes.
  7 Nightly Sync     v005 C1-C4 replay on the new copy (and the baseline, for reference).
  8 BO merge         every TableArchiveBO value reaches TableArchiveFRM10_12, or is a listed conflict.

Degrades gracefully: a table or the D2 columns not built yet is a SKIP/WARN line, not a crash.
--strict turns those into FAILs (the final go-live gate). Default output is failures only.

Reads xlsx with zipfile (table XML: name, ref, headers, GUID) + openpyxl read_only (cells, bounded
by the table ref). Never a full openpyxl load, never Excel.
"""
import argparse
import collections
import datetime as dt
import html
import json
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import mirror_lib as ML  # noqa: E402
from gen_data_inventory import table_columns  # noqa: E402

REPO = os.path.dirname(HERE)
CLIENT = os.path.dirname(REPO)
REFRESHER = os.path.join(REPO, "Office Scripts", "Mixed Query Refresher - Live Version.osts")
VIEWER_PQ = os.path.join(CLIENT, "FRM10-12", "viewer", "power-query")

FRM1012, FRM11, FRM13, BO = "TableArchiveFRM10_12", "TableArchiveFRM11", "TableArchiveFRM13", "TableArchiveBO"
TABLE_GUID = "{17014A43-5D94-47C6-982F-45D962DA4036}"
LIST_TABLES = collections.OrderedDict([
    ("TableArchiveOrderItems", "Order Items"),
    ("TableArchiveOrder", "Order"),
    ("TableArchiveModels", "Models"),
    ("TableArchiveModelRevisions", "Model Revisions"),
    ("TableArchiveClients", "Clients"),
])

# Pinned from workbooks/Archive active 2026-09-27 2356.xlsx (hard constraint 2).
PINNED_FRM1012 = [
    "Order", "Client", "KVA and KV", "Primary Voltage", "Secondary Voltage", "Phases", "JS #", "Description",
    "Type", "PO", "Order Date", "Lead Time", "Ing. Due Date", "Qty", "PO Item #", "Family", "Duplicate",
    "Engineering Required", "Duplicate Order", "Price", "Province/State", "WET-WETP", "Indexing", "LDs",
    "Initial Promised Date", "Trimestrial Customer", "Client Date Status", "Info+", "Protector Status",
    "Protector & Switchgear PO", "Protector & Switchgear Item #", "Sales Notes", "Technical Notes", "Location",
    "Status", "Witness/Other", "Temperature Rise", "Impulse", "DB", "Partial D", "Oil Analysis", "SFRA", "CSA",
    "Core", "Core Status", "Oil Type", "Oil Amount", "Production Line", "Configuration", "Section Qty", "Cable",
    "Coil Winder", "Form", "Copper (LV)", "Wire (HV)", "Overcoil", "Winder", "Time (days)", "Tank",
    "Tank Delivery Date", "Frame", "ISO Stack", "ISO Coil", "Lead Assembly", "Coiling Date", "Stacking Date",
    "Assembly Date", "Drying Date", "Tanking Date", "Testing Date", "Finishing Date", "Delivery Date",
    "Original Tanking Date", "Estimated Delivery Date", "Tanking date change justification",
    "Manual Estimated Delivery Date", "BO", "Price Value", "Price CAD", "Price USD", "Navigation Order",
    "Navigation Model", "Archived", "Lot", "Tanking Date Status", "Planning Notes", "Production Complexity",
    "__PowerAppsId__", "Client Desired Date", "FI", "Stack", "Production Status", "Last Synchronisation Date",
]
PINNED_COUNTS = {FRM1012: 93, FRM11: 39, FRM13: 55}

# D2: TableArchiveBO's detail headers (exact spelling, "Numbre" included) -> Order Items internal names.
BO_PARTS = [("Part Numbre", "PartNumber"), ("Description", "Description"), ("PO Intern", "POIntern"),
            ("Date", "Date"), ("Fournisseur Interne", "Fournisseur"), ("OK", "OK")]
BO_DETAIL = ["BO%d %s" % (i, h) for i in (1, 2, 3) for h, _ in BO_PARTS]
BO_OI_FIELD = dict([("BO", "BO")] + [("BO%d %s" % (i, h), "BO%d%s" % (i, f)) for i in (1, 2, 3) for h, f in BO_PARTS])

D3_COLS = {"Lead Time", "Price", "Estimated Delivery Date", "Price CAD", "Price USD"}   # recipe 12, 20, 74, 79, 80
E7_ORDERS = {"22140", "22141", "22156", "22157", "P00005", "P10003"}
STAGES = [("Coiling Date", "Coiling"), ("Stacking Date", "Stacking"), ("Assembly Date", "Assembly"),
          ("Drying Date", "Drying"), ("Testing Date", "Testing"), ("Finishing Date", "Finishing")]
STAGE_DATE_COLS = {c for c, _ in STAGES}
SYNC_COL = "Last Synchronisation Date"

READERS = [
    ("FRM11 Rows to purge", ["Order", "Location", "Status"]),
    ("FRM13", ["Order", "Location"]),
    ("viewer purge", ["Order", "Location", "Delivery Date"]),
    ("Query2", ["PO Item #", "Order", "Client", "Order Date"]),
    ("Power BI PriceReg", ["Price", "Price Value"]),
]
MIRROR_ADDED_OPTIONAL = re.compile(r"_Email$")   # mirror-added person e-mails: WARN if absent, not FAIL


# ======================================================================= shared vocabularies (read, never retyped)
def _read(path):
    with open(path, encoding="utf-8-sig") as f:
        return f.read()


def load_text_date_exempt(path=REFRESHER):
    m = re.search(r"TEXT_COLUMNS_WITH_DATE_IN_NAME\s*=\s*\[(.*?)\]", _read(path), re.S)
    if not m:
        raise SystemExit("ABORT: TEXT_COLUMNS_WITH_DATE_IN_NAME not found in %s" % path)
    return set(re.findall(r'"([^"]+)"', m.group(1)))


def load_location_codes(path=os.path.join(VIEWER_PQ, "LocationCodes.pq")):
    pairs = re.findall(r'Code\s*=\s*"([^"]+)"\s*,\s*Name\s*=\s*"([^"]+)"', _read(path))
    if len(pairs) != 12:
        raise SystemExit("ABORT: expected 12 Location pairs in %s, found %d" % (path, len(pairs)))
    return {name: code for code, name in pairs}


def load_stamp_codes(path=os.path.join(VIEWER_PQ, "StatusStampCodes.pq")):
    out = {"step": {}, "month": {}}
    for kind, key, code in re.findall(r'Kind\s*=\s*"(\w+)"\s*,\s*Key\s*=\s*"([^"]+)"\s*,\s*Code\s*=\s*"([^"]+)"', _read(path)):
        out[kind][key] = code
    if len(out["month"]) != 12 or not out["step"]:
        raise SystemExit("ABORT: StatusStampCodes.pq did not parse (%d steps, %d months)" % (len(out["step"]), len(out["month"])))
    return out


class Vocab:
    def __init__(self, exempt=None, locations=None, stamps=None):
        self.exempt = load_text_date_exempt() if exempt is None else exempt
        self.loc = load_location_codes() if locations is None else locations
        self.stamps = load_stamp_codes() if stamps is None else stamps

    def is_date_col(self, name):
        return "date" in name.lower() and name not in self.exempt


# ======================================================================= workbook model
def _col_num(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def parse_ref(ref):
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", ref or "")
    if not m:
        return None
    return _col_num(m.group(1)), int(m.group(2)), _col_num(m.group(3)), int(m.group(4))


class Table:
    """An Excel table: headers, rows (tuples, lazily read) and identity (GUID, xml id, part)."""

    def __init__(self, name, headers, rows=None, uid=None, xml_id=None, part=None, sheet=None, ref=None, loader=None):
        self.name, self.headers = name, list(headers)
        self._rows, self._loader = rows, loader
        self.uid, self.xml_id, self.part, self.sheet, self.ref = uid, xml_id, part, sheet, ref

    @property
    def rows(self):
        if self._rows is None:
            n = len(self.headers)
            self._rows = [tuple(r[:n]) + (None,) * (n - len(r)) for r in self._loader()]
        return self._rows

    def idx(self, h):
        return self.headers.index(h) if h in self.headers else None

    def keyed(self, key_col):
        """{key: row} (first wins) and the duplicate keys."""
        i = self.idx(key_col)
        out, dups = {}, []
        if i is None:
            return out, dups
        for r in self.rows:
            k = norm_key(r[i])
            if k is None:
                continue
            if k in out:
                dups.append(k)
            else:
                out[k] = r
        return out, dups


class Book:
    def __init__(self, tables, label=""):
        self.tables, self.label = tables, label

    def get(self, name):
        return self.tables.get(name)

    @classmethod
    def from_xlsx(cls, path):
        z = zipfile.ZipFile(path)
        names = set(z.namelist())
        wbx = z.read("xl/workbook.xml").decode("utf-8")
        rels = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', z.read("xl/_rels/workbook.xml.rels").decode("utf-8")))
        rels.update({a: b for b, a in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', z.read("xl/_rels/workbook.xml.rels").decode("utf-8"))})
        part_sheet = {}
        for tag in re.findall(r"<sheet\b[^>]*/>", wbx):
            sname = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'r:id="([^"]+)"', tag).group(1)
            target = rels[rid].lstrip("/")
            target = target if target.startswith("xl/") else "xl/" + target
            srels = target.replace("worksheets/", "worksheets/_rels/") + ".rels"
            if srels not in names:
                continue
            for t in re.findall(r'Target="([^"]*tables/table\d+\.xml)"', z.read(srels).decode("utf-8")):
                part_sheet["xl/tables/" + t.split("tables/")[-1]] = sname
        state = {"wb": None}

        def open_wb():
            if state["wb"] is None:
                import openpyxl
                state["wb"] = openpyxl.load_workbook(path, read_only=True, data_only=True)
            return state["wb"]

        tables = {}
        for part in sorted(n for n in names if re.match(r"xl/tables/table\d+\.xml$", n)):
            x = z.read(part).decode("utf-8", "ignore")
            head = re.search(r"<table\b[^>]*>", x).group(0)
            disp = re.search(r'\bdisplayName="([^"]+)"', head)
            name = html.unescape(disp.group(1)) if disp else table_columns(z, part)[0]
            _, _, cols = table_columns(z, part)
            cols = [html.unescape(c) for c in cols]
            uid = re.search(r'\bxr:uid="([^"]+)"', head)
            xid = re.search(r'\bid="([^"]+)"', head)
            ref = re.search(r'\bref="([^"]+)"', head).group(1)
            sheet = part_sheet.get(part)

            def loader(sheet=sheet, ref=ref):
                c0, r0, c1, r1 = parse_ref(ref)
                ws = open_wb()[sheet]
                if r1 <= r0:
                    return []
                return list(ws.iter_rows(min_row=r0 + 1, max_row=r1, min_col=c0, max_col=c1, values_only=True))
            tables[name] = Table(name, cols, uid=uid.group(1) if uid else None, xml_id=xid.group(1) if xid else None,
                                 part=part, sheet=sheet, ref=ref, loader=loader)
        return cls(tables, label=os.path.basename(path))


# ======================================================================= mirror model
class Mirror:
    """The live mirror's CSV output (list -> rows as dicts), plus its Columns catalog."""

    def __init__(self, tables, label=""):
        self.tables, self.label = tables, label
        self._by_id = {}
        self.catalog = ML.Catalog(self)

    def has(self, t):
        return t in self.tables

    def table(self, t):
        return self.tables[t]

    @classmethod
    def from_live(cls, folder=None):
        folder = folder or ML.LIVE
        tables = {}
        for t in list(LIST_TABLES.values()) + [ML.CATALOG]:
            p = os.path.join(folder, t + ".csv")
            if os.path.exists(p):
                tables[t] = ML.read_table(p)
        return cls(tables, label=folder)

    def columns(self, lst):
        rows = self.tables.get(lst) or []
        return list(rows[0].keys()) if rows else []

    def by_id(self, lst):
        if lst not in self._by_id:
            self._by_id[lst] = {ML.key(r): r for r in self.tables.get(lst, [])}
        return self._by_id[lst]

    def oi_by_title(self):
        if "_oi_title" not in self._by_id:
            self._by_id["_oi_title"] = {(r.get("Title") or "").strip(): r for r in self.tables.get("Order Items", []) if (r.get("Title") or "").strip()}
        return self._by_id["_oi_title"]

    def cat_rows(self, lst):
        return [r for r in self.catalog.rows if r["list"] == lst]

    def date_only_cols(self, lst):
        return {(r.get("csvColumn") or "").strip() or r["internalName"] for r in self.cat_rows(lst)
                if r.get("type") == "dateTime (dateOnly)"}

    def lookup_display_cols(self, lst):
        present = set(self.columns(lst))
        return {r["internalName"] for r in self.cat_rows(lst)
                if (r.get("type") or "").startswith("lookup") and r["internalName"] in present}


# ======================================================================= value normalisation
US_DT = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4}) (\d{1,2}):(\d{2})(?::(\d{2}))? ?([AP]M)$", re.I)
ISO_DT = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ](\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.\d+)?Z?)?$")
NUM = re.compile(r"^-?\d+(?:[.,]\d+)?$")


def norm_key(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _dt_canon(d):
    if isinstance(d, dt.datetime):
        if (d.hour, d.minute, d.second, d.microsecond) == (0, 0, 0, 0):
            return ("D", d.date())
        return ("DT", d.replace(microsecond=0))
    return ("D", d)


def parse_date_text(s):
    m = ISO_DT.match(s)
    try:
        if m:
            y, mo, d, hh, mi, ss = m.groups()
            return dt.datetime(int(y), int(mo), int(d), int(hh or 0), int(mi or 0), int(ss or 0))
        m = US_DT.match(s)
        if m:
            mo, d, y, hh, mi, ss, ap = m.groups()
            hh = int(hh) % 12 + (12 if ap.upper() == "PM" else 0)
            return dt.datetime(int(y), int(mo), int(d), hh, int(mi), int(ss or 0))
    except ValueError:
        return None
    return None


def canon(v):
    """A type-blind comparable form: '' / 'null' == None, '3' == 3, '24,9' == 24.9, text date == date."""
    if v is None:
        return None
    if isinstance(v, bool):
        return ("B", v)
    if isinstance(v, (int, float)):
        return ("N", round(float(v), 6))
    if isinstance(v, (dt.datetime, dt.date)):
        return _dt_canon(v)
    s = str(v).strip()
    if s == "" or s.lower() == "null":
        return None
    if s.lower() in ("true", "false"):
        return ("B", s.lower() == "true")
    if NUM.match(s):
        return ("N", round(float(s.replace(",", ".")), 6))
    d = parse_date_text(s)
    if d is not None:
        return _dt_canon(d)
    return ("S", s)


def identical(a, b):
    return type(a) is type(b) and a == b


def is_blank(v):
    return canon(v) is None


def sp_date(raw):
    """A mirror date-only raw value -> its stored date (first 10 chars; never shifted)."""
    s = (raw or "").strip() if isinstance(raw, str) else raw
    if not s:
        return None
    if isinstance(s, (dt.datetime, dt.date)):
        return s.date() if isinstance(s, dt.datetime) else s
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return None


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def show(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour, v.minute) != (0, 0) else v.strftime("%Y-%m-%d")
    return repr(v)


# ======================================================================= the recipe (docs/archive-frm10-12-column-recipe.md)
def _txt(v):
    return v if (v is not None and str(v).strip() != "") else None


def _num(v):
    s = "" if v is None else str(v).strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return s


def _num0(v):
    n = _num(v)
    return 0.0 if n is None else n


def _mark(ch):
    return lambda v: ch if str(v).strip().lower() == "true" else None


def _bool(v):
    s = str(v).strip().lower()
    return True if s == "true" else False if s == "false" else None


def _multi(v):
    s = (v or "").strip()
    if not s:
        return None
    try:
        vals = json.loads(s)
    except ValueError:
        return s
    if isinstance(vals, list):
        return ", ".join(str(x) for x in vals) or None
    return str(vals)


SIMPLE = {   # column -> (source list, internal field, transform)
    "Order": ("OI", "Title", _txt), "Client": ("ORD", "Client", _txt),
    "KVA and KV": ("MR", "kVA_x0020_and_x0020_kV", _txt), "Primary Voltage": ("MR", "PrimaryVoltage", _num),
    "Secondary Voltage": ("MR", "SecondaryVoltage", _num), "Phases": ("MR", "Phases", _num),
    "JS #": ("MR", "JS_x0020__x0023_", _txt), "Description": ("MR", "Description", _multi),
    "Type": ("MR", "Model_x0020_Type", _txt), "PO": ("ORD", "PO", _txt), "Order Date": ("ORD", "Order_x0020_Date", sp_date),
    "Qty": ("OI", "Qty", _num), "PO Item #": ("MR", "ModelName", _txt), "Family": ("MR", "Family", _txt),
    "Engineering Required": ("ORD", "EngineeringRequired", _mark("Y")), "Price": ("ORD", "Price", _num0),
    "Province/State": ("ORD", "Province_x002F_State", _txt), "WET-WETP": ("ORD", "WET_x002d_WETP", _txt),
    "Indexing": ("ORD", "Indexing", _txt), "LDs": ("ORD", "LDs", _mark("Y")),
    "Initial Promised Date": ("ORD", "Initial_x0020_Promised_x0020_Dat", sp_date),
    "Trimestrial Customer": ("OI", "TrimestrialCustomer", _txt), "Client Date Status": ("ORD", "ClientDateStatus", _txt),
    "Info+": ("OI", "Info_x002b_", _txt), "Protector Status": ("OI", "ProtectorStatus", _txt),
    "Protector & Switchgear PO": ("OI", "ProtectorSwitchgearPO", _txt),
    "Protector & Switchgear Item #": ("OI", "Protector_x0020__x0026__x0020_Sw", _txt),
    "Sales Notes": ("ORD", "SalesNotes", _txt), "Technical Notes": ("OI", "Technical_x0020_Notes", _txt),
    "Witness/Other": ("OI", "Witness_x002f_Other", _txt),
    "Temperature Rise": ("OI", "TemperatureRise", _mark("x")), "Impulse": ("OI", "Impulse", _mark("x")),
    "DB": ("OI", "DB", _mark("x")), "Partial D": ("OI", "PartialD", _mark("x")), "Oil Analysis": ("OI", "OilAnalysis", _mark("x")),
    "SFRA": ("OI", "SFRA", _mark("Y")), "CSA": ("OI", "CSA", _bool), "Core": ("MR", "Core_x0020_Type", _txt),
    "Core Status": ("OI", "CoreStatus", _txt), "Oil Type": ("MR", "Oil_x0020_Type", _txt), "Oil Amount": ("MR", "OilAmount", _num),
    "Production Line": ("OI", "ProductionLine", _txt), "Configuration": ("OI", "Configuration", _txt),
    "Section Qty": ("OI", "Section_x0020_Qty", _num), "Cable": ("MR", "Cable", _txt), "Coil Winder": ("OI", "CoilWinder", _txt),
    "Form": ("MR", "Form", _txt), "Copper (LV)": ("MR", "Copper_x0028_LV_x0029_", _txt), "Wire (HV)": ("MR", "Wire_x0028_HV_x0029_", _txt),
    "Overcoil": ("MR", "Overcoil", _num), "Winder": ("OI", "Winder", _txt), "Time (days)": ("OI", "Time_x0028_days_x0029_", _num),
    "Tank": ("OI", "Tank", _mark("R")), "Tank Delivery Date": ("OI", "TankDeliveryDate", sp_date), "Frame": ("OI", "Frame", _txt),
    "ISO Stack": ("OI", "ISOStack", _mark("R")), "ISO Coil": ("OI", "ISOCoil", _mark("R")), "Lead Assembly": ("OI", "LeadAssembly", _mark("R")),
    "Tanking Date": ("OI", "Planned_x0020_Tanking_x0020_Date", sp_date),
    "Delivery Date": ("OI", "Planned_x0020_Delivery_x0020_Dat", sp_date),
    "Original Tanking Date": ("OI", "OriginalTankingDate", sp_date),
    "Estimated Delivery Date": ("OI", "EstimatedDeliveryDate", sp_date),
    "Tanking date change justification": ("OI", "TankingDateChangeJustification", _txt),
    "Manual Estimated Delivery Date": ("OI", "ManualEstimatedDeliveryDate", sp_date),
    "BO": ("OI", "BO", _txt), "Price Value": ("ORD", "Price", _num),
    "Price CAD": ("OI", "PriceCAD", _num), "Price USD": ("OI", "PriceUSD", _num),
}
SRC_LIST = {"OI": "Order Items", "ORD": "Order", "MR": "Model Revisions", "CLI": "Clients"}
OI_FIELDS_USED = sorted({f for s, f, _ in SIMPLE.values() if s == "OI"} | {"TankingDate", "StepStatus", "StatusDate", "Status", "Location"}
                        | {"%sDate" % s for _, s in STAGES} | {"%sStatus" % s for _, s in STAGES})


def context(mirror, oi):
    ordr = mirror.by_id("Order").get(ML.key({"Id": oi.get("OrderNumberId")}), {}) if oi.get("OrderNumberId") else {}
    mr = mirror.by_id("Model Revisions").get(ML.key({"Id": oi.get("ModelRevisionId")}), {}) if oi.get("ModelRevisionId") else {}
    cli = mirror.by_id("Clients").get(ML.key({"Id": ordr.get("ClientId")}), {}) if ordr.get("ClientId") else {}
    return {"OI": oi, "ORD": ordr, "MR": mr, "CLI": cli}


def expected(col, ctx, vocab):
    """(computable, value) for an FRM10-12 column from the mirror, per the recipe."""
    oi, ordr, mr, cli = ctx["OI"], ctx["ORD"], ctx["MR"], ctx["CLI"]
    if col in SIMPLE:
        src, field, fn = SIMPLE[col]
        return True, fn(ctx[src].get(field))
    if col == "Location":
        v = _txt(oi.get("Location"))
        return True, vocab.loc.get(v, v) if v else None
    lead = _num(cli.get("CliLeadTimeWeeks"))
    lead = lead if isinstance(lead, float) else 26.0
    if col == "Lead Time":
        return True, lead
    if col == "Ing. Due Date":
        base = sp_date(oi.get("TankingDate")) or sp_date(ordr.get("Initial_x0020_Promised_x0020_Dat"))
        return True, (base - dt.timedelta(days=int((lead + 4) * 7))) if base else None
    if col == "Status":
        step = vocab.stamps["step"].get(_txt(oi.get("StepStatus")) or "")
        when = sp_date(oi.get("StatusDate"))
        mon = vocab.stamps["month"].get(str(when.month)) if when else None
        if step is None or mon is None:
            return True, _txt(oi.get("Status"))
        return True, "%s-%s-%d" % (step, mon, when.day)
    for c, stage in STAGES:
        if col == c:
            if (oi.get(stage + "Status") or "").strip() == "In Progress":
                return True, "EC"
            return True, sp_date(oi.get(stage + "Date"))
    if col == "Navigation Order":
        o = _txt(oi.get("Title"))
        return True, ("Ouvrir commande " + o.split("-", 1)[0]) if o and "-" in o else None
    if col == "Navigation Model":
        m = _txt(mr.get("ModelName"))
        return True, ("Ouvrir modèle " + m) if m else None
    if col in ("Duplicate", "Duplicate Order"):
        return True, None
    return False, None   # legacy / local-only columns and the sync stamp: no SharePoint source


# ======================================================================= results
LEVELS = ["PASS", "SKIP", "WARN", "FAIL"]


class Result:
    def __init__(self, num, title):
        self.num, self.title, self.items, self.detail = num, title, [], []

    def add(self, level, msg, examples=None):
        self.items.append((level, msg, list(examples or [])[:5]))

    @property
    def status(self):
        lv = {i[0] for i in self.items}
        for s in ("FAIL", "WARN"):
            if s in lv:
                return s
        if lv and lv <= {"SKIP"}:
            return "SKIP"
        return "PASS"

    def has_fail(self, text=None):
        return any(l == "FAIL" and (text is None or text in m) for l, m, _ in self.items)


def missing(res, what, strict):
    res.add("FAIL" if strict else "SKIP", "%s not present: %s" % (what, "required by --strict" if strict else "skipped"))


# ======================================================================= check 1: shape
def check_shape(base, new, opts):
    res = Result(1, "Shape")
    for name in (FRM11, FRM13):
        b, n = base.get(name), new.get(name)
        if b is None:
            res.add("FAIL", "%s missing from the baseline" % name)
            continue
        if len(b.headers) != PINNED_COUNTS[name]:
            res.add("FAIL", "baseline %s has %d columns, pinned %d: wrong baseline file?" % (name, len(b.headers), PINNED_COUNTS[name]))
        if n is None:
            res.add("FAIL", "%s missing from the new copy" % name)
            continue
        if n.headers != b.headers:
            res.add("FAIL", "%s headers changed" % name, header_diff(b.headers, n.headers))
            continue
        diffs = []
        if len(n.rows) != len(b.rows):
            diffs.append("row count %d -> %d" % (len(b.rows), len(n.rows)))
        for i, (rb, rn) in enumerate(zip(b.rows, n.rows)):
            if rb != rn:
                cols = [h for h, x, y in zip(b.headers, rb, rn) if x != y]
                diffs.append("row %d (%s): %s" % (i + 2, rb[0], ", ".join("%s %s->%s" % (h, show(rb[b.idx(h)]), show(rn[n.idx(h)])) for h in cols[:3])))
        if diffs:
            res.add("FAIL", "%s cells changed (%d rows differ); it must stay untouched" % (name, len(diffs)), diffs)
        else:
            res.add("PASS", "%s: %d columns, %d rows, identical" % (name, len(n.headers), len(n.rows)))
        if n.ref != b.ref:
            res.add("WARN", "%s ref %s -> %s although the cells are identical" % (name, b.ref, n.ref))

    b, n = base.get(FRM1012), new.get(FRM1012)
    if b is None or n is None:
        res.add("FAIL", "%s missing from the %s" % (FRM1012, "baseline" if b is None else "new copy"))
        return res
    if b.headers != PINNED_FRM1012:
        res.add("FAIL", "baseline %s is not the pinned 93 (09-27 23:56 shape)" % FRM1012, header_diff(PINNED_FRM1012, b.headers))
    want = PINNED_FRM1012 + BO_DETAIL
    if n.headers == want:
        res.add("PASS", "%s headers = the pinned 93 then the 18 BO detail columns" % FRM1012)
    elif n.headers == PINNED_FRM1012:
        res.add("FAIL" if opts.strict else "WARN", "%s has the 93 but not yet the 18 BO detail columns (D2)%s"
                % (FRM1012, "" if opts.strict else "; BO-column checks skipped"))
    else:
        res.add("FAIL", "%s headers are not the pinned 93 + 18 BO detail columns" % FRM1012, header_diff(want, n.headers))
    # Identity. The Excel Online (Graph) connector's table id is the table part's xr:uid GUID.
    if n.uid != TABLE_GUID:
        res.add("FAIL", "%s table GUID (xr:uid) is %s, must be %s: table was recreated" % (FRM1012, n.uid, TABLE_GUID))
    else:
        res.add("PASS", "%s xr:uid = %s" % (FRM1012, TABLE_GUID))
    if b.uid != TABLE_GUID:
        res.add("WARN", "baseline %s GUID is %s, expected %s" % (FRM1012, b.uid, TABLE_GUID))
    if (n.xml_id, n.part, n.sheet) != (b.xml_id, b.part, b.sheet):
        res.add("WARN", "%s table xml id/part/sheet moved: %s -> %s (the GUID is what the connector uses)"
                % (FRM1012, (b.xml_id, b.part, b.sheet), (n.xml_id, n.part, n.sheet)))
    # Keys
    i = n.idx("Order")
    if i is not None:
        keys = [norm_key(r[i]) for r in n.rows]
        dups = [k for k, c in collections.Counter(k for k in keys if k).items() if c > 1]
        if dups:
            res.add("FAIL", "%s has %d duplicate Order keys" % (FRM1012, len(dups)), dups)
        blanks = sum(1 for k in keys if k is None)
        bi = b.idx("Order")
        bblanks = sum(1 for r in b.rows if norm_key(r[bi]) is None) if bi is not None else 0
        if blanks > bblanks:
            res.add("FAIL", "%s has %d rows with a blank Order (baseline %d)" % (FRM1012, blanks, bblanks))
        elif blanks:
            res.add("WARN", "%s keeps %d blank-Order row(s) carried from the baseline" % (FRM1012, blanks))
    return res


def header_diff(want, got):
    out = []
    ws, gs = set(want), set(got)
    for h in want:
        if h not in gs:
            out.append("missing %r" % h)
    for h in got:
        if h not in ws:
            out.append("extra %r" % h)
    for k, (a, c) in enumerate(zip(want, got)):
        if a != c:
            out.append("position %d: want %r, got %r" % (k + 1, a, c))
            break
    if len(want) != len(got):
        out.append("count %d, want %d" % (len(got), len(want)))
    return out


# ======================================================================= check 2: FRM10-12 cells
CLASSES = ["identical", "cleaner", "newer", "intended", "REGRESSION"]


def classify(col, bv, nv, ctx, vocab, order, bo_value=None):
    """-> (class, stale) ; stale = identical to the baseline but the mirror says otherwise (info only).
    bo_value: the baseline TableArchiveBO value of this column for this Order (D2 backfill), if any."""
    ok, exp = expected(col, ctx, vocab) if ctx else (False, None)
    if identical(bv, nv):
        return "identical", bool(ok and canon(nv) != canon(exp))
    cb, cn = canon(bv), canon(nv)
    worse_type = isinstance(nv, str) and bv is not None and not isinstance(bv, str)
    if cb == cn and not worse_type:
        return "cleaner", False
    if col == SYNC_COL:   # the archive's own stamp: may only move forward
        fwd = isinstance(nv, (dt.datetime, dt.date)) and (as_date(bv) is None or as_date(nv) >= as_date(bv))
        return ("newer" if fwd else "REGRESSION"), False
    matches_mirror = ok and cn == canon(exp) and not (isinstance(nv, str) and isinstance(exp, (dt.date, float, bool)))
    if col == "BO" and cb is None and bo_value is not None and cn == canon(bo_value) and not matches_mirror:
        return "intended", False   # D2: historical blank filled from TableArchiveBO
    e7 = order.split("-", 1)[0] in E7_ORDERS and vocab.is_date_col(col) and col in SIMPLE
    if col in D3_COLS or e7:
        if matches_mirror:
            return "intended", False
        return "REGRESSION", False
    if matches_mirror:
        return "newer", False
    return "REGRESSION", False


def check_cells(base, new, mirror, vocab, opts):
    res = Result(2, "FRM10-12 cells")
    b, n = base.get(FRM1012), new.get(FRM1012)
    if b is None or n is None:
        res.add("SKIP", "%s missing from a copy: skipped" % FRM1012)
        return res
    bad_src = [f for f in OI_FIELDS_USED if mirror.catalog.synced("Order Items", f)]
    if bad_src:
        res.add("FAIL", "the recipe sources %d FRM10-12 value(s) from Order Items sync copies" % len(bad_src), bad_src)
    bk, _ = b.keyed("Order")
    nk, _ = n.keyed("Order")
    bot = base.get(BO)
    bo_bo = {}
    if bot is not None and "BO" in bot.headers:
        bo_bo = {k: r[bot.idx("BO")] for k, r in bot.keyed("Order")[0].items()}
    oi_t = mirror.oi_by_title()
    cols = [h for h in PINNED_FRM1012 if h in b.headers and h in n.headers]
    counts = collections.defaultdict(collections.Counter)
    stale = collections.Counter()
    ex = collections.defaultdict(list)
    both = [k for k in bk if k in nk]
    for k in both:
        rb, rn = bk[k], nk[k]
        oi = oi_t.get(k)
        ctx = context(mirror, oi) if oi else None
        for c in cols:
            bv, nv = rb[b.idx(c)], rn[n.idx(c)]
            cls, st = classify(c, bv, nv, ctx, vocab, k, bo_bo.get(k) if c == "BO" else None)
            counts[c][cls] += 1
            if st:
                stale[c] += 1
            if cls != "identical" and len(ex[(c, cls)]) < 5:
                e = "%s: %s -> %s" % (k, show(bv), show(nv))
                if cls == "REGRESSION":
                    ok, exp = expected(c, ctx, vocab) if ctx else (False, None)
                    e += "  (mirror: %s)" % (show(exp) if ok else ("unit not on SharePoint" if not ctx else "no source"))
                ex[(c, cls)].append(e)
    tot = collections.Counter()
    for c in cols:
        tot.update(counts[c])
    res.add("PASS" if not tot["REGRESSION"] else "FAIL",
            "%d Orders in both copies, %d cells: %s" % (len(both), sum(tot.values()), ", ".join("%s %d" % (k, tot[k]) for k in CLASSES)))
    for c in cols:
        if counts[c]["REGRESSION"]:
            res.add("FAIL", "%s: %d REGRESSION" % (c, counts[c]["REGRESSION"]), ex[(c, "REGRESSION")])
    res.detail.append("| column | " + " | ".join(CLASSES) + " | stale* |")
    res.detail.append("|---|" + "---:|" * (len(CLASSES) + 1))
    for c in cols:
        res.detail.append("| %s | %s | %d |" % (c, " | ".join(str(counts[c][k]) for k in CLASSES), stale[c]))
    res.detail.append("")
    res.detail.append("*stale = identical to the baseline but differs from the mirror's value (info: the row was not refreshed).")
    for (c, cls), es in sorted(ex.items()):
        if cls in ("cleaner", "newer", "intended"):
            res.detail.append("- %s / %s: %s" % (c, cls, "; ".join(es)))
    if sum(stale.values()):
        res.add("PASS", "info: %d cells identical to the baseline but not to the mirror (stale rows)" % sum(stale.values()))
    return res


# ======================================================================= check 3: no text in date columns
def check_text_dates(base, new, mirror, vocab, opts):
    res = Result(3, "Date columns hold no text")
    checked = 0
    for name, t in sorted(new.tables.items()):
        if not name.startswith("TableArchive"):
            continue
        if name in LIST_TABLES:
            date_cols = [h for h in t.headers if h in mirror.date_only_cols(LIST_TABLES[name])]
            key = "Id"
        else:
            date_cols = [h for h in t.headers if vocab.is_date_col(h)]
            key = "Order" if "Order" in t.headers else t.headers[0]
        if not date_cols:
            continue
        bt = base.get(name)
        legacy = set()
        if bt is not None and key in bt.headers:
            ki = bt.idx(key)
            for c in date_cols:
                ci = bt.idx(c)
                if ci is None:
                    continue
                for r in bt.rows:
                    if isinstance(r[ci], str):
                        legacy.add((norm_key(r[ki]), c, r[ci]))
        ki = t.idx(key)
        bad, old = [], []
        for c in date_cols:
            ci = t.idx(c)
            for r in t.rows:
                v = r[ci]
                checked += 1
                if not isinstance(v, str) or v.strip() == "":
                    continue
                if c in STAGE_DATE_COLS and name == FRM1012 and v.strip() == "EC":
                    continue
                if (norm_key(r[ki]) if ki is not None else None, c, v) in legacy:
                    old.append("%s %s=%r" % (norm_key(r[ki]), c, v))
                else:
                    bad.append("%s %s=%r" % (norm_key(r[ki]) if ki is not None else "?", c, v))
        if bad:
            res.add("FAIL", "%s: %d text value(s) in date columns" % (name, len(bad)), bad)
        if old:
            res.add("WARN", "%s: %d legacy text value(s) kept exactly as in the baseline" % (name, len(old)), old)
        if not bad and not old:
            res.add("PASS", "%s: %d date columns clean" % (name, len(date_cols)))
    if not checked:
        res.add("SKIP", "no archive date columns found")
    return res


# ======================================================================= check 4: the five list tables vs the mirror
def check_list_tables(base, new, mirror, opts):
    res = Result(4, "List tables vs mirror")
    for tname, lst in LIST_TABLES.items():
        t = new.get(tname)
        if t is None:
            missing(res, tname, opts.strict)
            continue
        if not mirror.has(lst):
            res.add("FAIL", "mirror has no live %s.csv" % lst)
            continue
        cat = mirror.catalog
        mcols = mirror.columns(lst)
        hdr = set(t.headers)
        miss, miss_opt = [], []
        for c in mcols:
            if c in hdr or cat.rest_field(lst, c) in hdr:
                continue
            (miss_opt if MIRROR_ADDED_OPTIONAL.search(c) else miss).append(c)
        if miss:
            res.add("FAIL", "%s lacks %d mirror column(s) (new SharePoint columns must be added)" % (tname, len(miss)), miss)
        if miss_opt:
            res.add("WARN", "%s lacks %d mirror-added column(s)" % (tname, len(miss_opt)), miss_opt)
        if "Id" not in hdr:
            res.add("FAIL", "%s has no Id column" % tname)
            continue
        keyed, dups = t.keyed("Id")
        if dups:
            res.add("FAIL", "%s has %d duplicate Id(s)" % (tname, len(dups)), dups)
        m = mirror.by_id(lst)
        lost = sorted(set(m) - set(keyed), key=_sortkey)
        extra = sorted(set(keyed) - set(m), key=_sortkey)
        if lost:
            res.add("FAIL", "%s is missing %d mirror Id(s)" % (tname, len(lost)), lost)
        if extra:
            first = opts.first_build or base.get(tname) is None
            res.add("FAIL" if first else "WARN", "%s has %d Id(s) no longer on SharePoint%s" % (
                tname, len(extra), " (first build: must equal the mirror)" if first else " (deleted items kept, D4)"), extra)
        if not lost and not dups and not (extra and (opts.first_build or base.get(tname) is None)):
            res.add("PASS", "%s: %d Ids match the mirror's %d" % (tname, len(keyed), len(m)))
        # values
        dcols = [c for c in mirror.date_only_cols(lst) if c in hdr]
        lcols = [c for c in mirror.lookup_display_cols(lst) if c in hdr]
        dbad, lbad = [], []
        for k, mr in m.items():
            r = keyed.get(k)
            if r is None:
                continue
            for c in dcols:
                want, got = sp_date(mr.get(c)), r[t.idx(c)]
                gd = got.date() if isinstance(got, dt.datetime) else got if isinstance(got, dt.date) else None
                if want != gd and not (want is None and is_blank(got)):
                    dbad.append("Id %s %s: archive %s, mirror %r" % (k, c, show(got), mr.get(c)))
            for c in lcols:
                got = r[t.idx(c)]
                if ML.norm(mr.get(c)) != ML.norm("" if got is None else got):
                    lbad.append("Id %s %s: archive %s, mirror %r" % (k, c, show(got), mr.get(c)))
        if dbad:
            res.add("FAIL", "%s: %d date-only value(s) differ from the mirror's first 10 chars" % (tname, len(dbad)), dbad)
        if lbad:
            res.add("FAIL", "%s: %d lookup display value(s) differ from the mirror" % (tname, len(lbad)), lbad)
        if not dbad and not lbad:
            res.add("PASS", "%s: %d date-only + %d lookup columns equal the mirror" % (tname, len(dcols), len(lcols)))
    return res


def _sortkey(k):
    return (0, int(k)) if str(k).isdigit() else (1, str(k))


# ======================================================================= check 5: additivity
KEYS = {FRM1012: "Order", FRM13: "Order", BO: "Order"}


def table_key(name, t):
    if name in LIST_TABLES:
        return "Id"
    return KEYS.get(name) or t.headers[0]


def check_additivity(base, new, opts):
    res = Result(5, "Additivity")
    for name, b in sorted(base.tables.items()):
        if not name.startswith("TableArchive"):
            continue
        n = new.get(name)
        if n is None:
            if name == BO:
                res.add("WARN", "%s removed from the new copy (D2 allows it only after one verified refresh cycle)" % name)
            else:
                res.add("FAIL", "%s is in the baseline but missing from the new copy" % name)
            continue
        key = table_key(name, b)
        if len(n.rows) < len(b.rows):
            res.add("FAIL", "%s row count dropped %d -> %d" % (name, len(b.rows), len(n.rows)))
        if key not in n.headers or key not in b.headers:
            res.add("FAIL", "%s: key column %r missing" % (name, key))
            continue
        bi, ni = b.idx(key), n.idx(key)
        nkeys = collections.Counter(norm_key(r[ni]) for r in n.rows)
        bkeys = collections.Counter(norm_key(r[bi]) for r in b.rows)
        lost = sorted((k for k, c in bkeys.items() if k is not None and nkeys.get(k, 0) < min(c, 1)), key=_sortkey)
        if lost:
            res.add("FAIL", "%s lost %d baseline key(s) (%s)" % (name, len(lost), key), lost)
        elif len(n.rows) >= len(b.rows):
            res.add("PASS", "%s: %d -> %d rows, every baseline %s kept" % (name, len(b.rows), len(n.rows), key))
    return res


# ======================================================================= check 6: readers
def check_readers(base, new, vocab, opts):
    res = Result(6, "Readers")
    n, b = new.get(FRM1012), base.get(FRM1012)
    if n is None:
        res.add("FAIL", "%s missing" % FRM1012)
        return res
    for who, cols in READERS:
        miss = [c for c in cols if c not in n.headers]
        if miss:
            res.add("FAIL", "%s reads %s: missing from %s" % (who, ", ".join(miss), FRM1012))
        else:
            res.add("PASS", "%s: %s present" % (who, ", ".join(cols)))
    li = n.idx("Location")
    if li is not None:
        legacy = {r[b.idx("Location")] for r in b.rows} if b is not None and "Location" in b.headers else set()
        allowed = set(vocab.loc.values()) | {v for v in legacy if isinstance(v, str)}
        bad = collections.Counter(r[li] for r in n.rows if not is_blank(r[li]) and r[li] not in allowed)
        if bad:
            res.add("FAIL", "Location holds %d value(s) that are not two-letter codes" % sum(bad.values()),
                    ["%r x%d" % kv for kv in bad.most_common(5)])
        else:
            res.add("PASS", "Location values are codes (12 from LocationCodes.pq + baseline legacy %s)"
                    % sorted(v for v in legacy if isinstance(v, str) and v and v not in vocab.loc.values()))
    si = n.idx("Status")
    if si is not None:
        pre = lambda v: str(v).split("-", 1)[0].strip()
        legacy = {pre(r[b.idx("Status")]) for r in b.rows if not is_blank(r[b.idx("Status")])} if b is not None and "Status" in b.headers else set()
        allowed = set(vocab.stamps["step"].values()) | legacy
        bad = collections.Counter(pre(r[si]) for r in n.rows if not is_blank(r[si]) and pre(r[si]) not in allowed)
        if bad:
            res.add("FAIL", "Status holds %d value(s) whose prefix is not a step code" % sum(bad.values()),
                    ["%r x%d" % kv for kv in bad.most_common(5)])
        else:
            res.add("PASS", "Status prefixes are step codes")
    return res


# ======================================================================= check 7: Nightly Sync replay (v005 C1-C4)
def eastern_today():
    utc = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
    local = utc + dt.timedelta(hours=-5)
    local = utc + dt.timedelta(hours=ML._eastern_offset_hours(local))
    return local.date()


def replay(table, mirror, cutoff):
    cand = {(u.get("Title") or "").strip(): u for u in mirror.tables.get("Order Items", [])
            if u.get("Location") == "Livraison" and u.get("ItemStatus") == "Delivered"}
    io, il, idd = table.idx("Order"), table.idx("Location"), table.idx("Delivery Date")
    li_old = {norm_key(r[io]) for r in table.rows if r[il] == "LI" and as_date(r[idd]) and as_date(r[idd]) <= cutoff}
    return cand, sorted(t for t in cand if t in li_old)


def check_nightly(base, new, mirror, opts):
    res = Result(7, "Nightly Sync replay")
    n = new.get(FRM1012)
    if n is None or any(n.idx(c) is None for c in ("Order", "Location", "Delivery Date")):
        res.add("FAIL", "%s or its Order/Location/Delivery Date columns missing: the gate cannot run" % FRM1012)
        return res
    today = opts.today or eastern_today()
    cutoff = today - dt.timedelta(days=7)
    cand, conf = replay(n, mirror, cutoff)
    oi = mirror.oi_by_title()
    wrong = [t for t in conf if not (oi.get(t, {}).get("Location") == "Livraison" and oi.get(t, {}).get("ItemStatus") == "Delivered")]
    if wrong:
        res.add("FAIL", "%d confirmed unit(s) not at Livraison + Delivered on the mirror" % len(wrong), wrong)
    b = base.get(FRM1012)
    bconf = replay(b, mirror, cutoff)[1] if b is not None else []
    res.add("PASS", "today %s, cutoff %s: %d candidates, %d confirmed on the new copy (baseline replay: %d)"
            % (today, cutoff, len(cand), len(conf), len(bconf)))
    gained, lost = sorted(set(conf) - set(bconf)), sorted(set(bconf) - set(conf))
    if gained:
        res.add("WARN", "%d unit(s) confirmed now but not on the baseline replay: review before deletes" % len(gained), gained)
    if lost:
        res.add("WARN", "%d unit(s) confirmed on the baseline replay but held back now (safe direction)" % len(lost), lost)
    res.detail.append("confirmed: " + ", ".join(conf))
    return res


# ======================================================================= check 8: BO merge
def check_bo(base, new, mirror, opts):
    res = Result(8, "BO merge")
    bo, n, b12 = base.get(BO), new.get(FRM1012), base.get(FRM1012)
    if bo is None:
        res.add("SKIP", "baseline has no %s: skipped" % BO)
        return res
    if n is None:
        res.add("FAIL", "%s missing from the new copy" % FRM1012)
        return res
    cols = [c for c in ["BO"] + BO_DETAIL if c in bo.headers and c in n.headers]
    absent = [c for c in BO_DETAIL if c not in n.headers]
    pre_d2 = bool(absent) and not opts.strict
    if absent:
        (res.add("FAIL", "%d BO detail column(s) absent from %s" % (len(absent), FRM1012), absent) if opts.strict else
         res.add("SKIP", "%s has no BO detail columns yet (D2 not applied): only BO checked, losses reported as WARN" % FRM1012))
    nk, _ = n.keyed("Order")
    bk12, _ = b12.keyed("Order") if b12 is not None else ({}, [])
    oi_t = mirror.oi_by_title()
    date_cols = mirror.date_only_cols("Order Items")
    lostv, oi_conf, old_conf, ok = [], [], [], 0
    boi = bo.idx("Order")
    for r in bo.rows:
        k = norm_key(r[boi])
        if k is None:
            continue
        for c in cols:
            v = r[bo.idx(c)]
            if is_blank(v) or v is False:   # BO / BOn OK: False is the "no BO" default, not a value
                continue
            row = nk.get(k)
            if row is None:
                lostv.append("%s %s=%s (Order not in %s)" % (k, c, show(v), FRM1012))
                continue
            nv = row[n.idx(c)]
            if canon(nv) == canon(v) or (as_date(nv) and as_date(nv) == as_date(v)):
                ok += 1
                continue
            f = BO_OI_FIELD[c]
            oiv = (oi_t.get(k) or {}).get(f)
            oiv = sp_date(oiv) if f in date_cols else oiv
            if not is_blank(oiv) and canon(oiv) != canon(v) and (canon(nv) == canon(oiv) or as_date(nv) == as_date(oiv)):
                oi_conf.append("%s %s: Order Items %s, TableArchiveBO %s" % (k, c, show(oiv), show(v)))
                continue
            old = bk12.get(k)
            if old is not None and c in b12.headers and not is_blank(old[b12.idx(c)]) and identical(old[b12.idx(c)], nv):
                old_conf.append("%s %s: archive kept %s, TableArchiveBO %s" % (k, c, show(nv), show(v)))
                continue
            lostv.append("%s %s: TableArchiveBO %s, new %s" % (k, c, show(v), show(nv)))
    if lostv:
        if pre_d2:
            res.add("WARN", "%d TableArchiveBO value(s) not (yet) in %s: the D2 backfill must bring them" % (len(lostv), FRM1012), lostv)
        else:
            res.add("FAIL", "%d TableArchiveBO value(s) lost" % len(lostv), lostv)
    if oi_conf:
        res.add("WARN", "%d conflict(s): Order Items value wins over TableArchiveBO (list for the user)" % len(oi_conf), oi_conf)
    if old_conf:
        res.add("WARN", "%d conflict(s): the archive already held a different value (kept)" % len(old_conf), old_conf)
    res.add("PASS", "%d TableArchiveBO values present in %s (%s)" % (ok, FRM1012, ", ".join(cols)))
    res.detail += ["- " + x for x in oi_conf + old_conf]
    return res


# ======================================================================= driver
def run_checks(base, new, mirror, vocab, opts):
    return [
        check_shape(base, new, opts),
        check_cells(base, new, mirror, vocab, opts),
        check_text_dates(base, new, mirror, vocab, opts),
        check_list_tables(base, new, mirror, opts),
        check_additivity(base, new, opts),
        check_readers(base, new, vocab, opts),
        check_nightly(base, new, mirror, opts),
        check_bo(base, new, mirror, opts),
    ]


def render_console(results, verbose=False):
    out = []
    for r in results:
        out.append("%-4s [%d] %s" % (r.status, r.num, r.title))
        for level, msg, ex in r.items:
            if level == "PASS" and not verbose:
                continue
            out.append("      %-4s %s" % (level, msg))
            if level in ("FAIL", "WARN") or verbose:
                for e in ex:
                    out.append("             - %s" % e)
        if verbose:
            out += ["      " + d for d in r.detail]
    nf = sum(1 for r in results if r.status == "FAIL")
    out.append("")
    out.append("RESULT: %s (%d of %d checks failed)" % ("FAIL" if nf else "PASS", nf, len(results)))
    return "\n".join(out)


def render_report(results, base, new, mirror, opts):
    out = ["# Archive rebuild check", "",
           "- baseline: `%s`" % base.label, "- new: `%s`" % new.label, "- mirror: `%s`" % mirror.label,
           "- run: %s%s" % (dt.datetime.now().strftime("%Y-%m-%d %H:%M"), " (strict)" if opts.strict else ""), "",
           "| # | check | status |", "|---|---|---|"]
    out += ["| %d | %s | %s |" % (r.num, r.title, r.status) for r in results]
    for r in results:
        out += ["", "## %d. %s: %s" % (r.num, r.title, r.status), ""]
        for level, msg, ex in r.items:
            out.append("- **%s** %s" % (level, msg))
            out += ["  - `%s`" % e for e in ex]
        if r.detail:
            out += [""] + r.detail
    return "\n".join(out) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--base", required=True)
    ap.add_argument("--new", required=True)
    ap.add_argument("--report")
    ap.add_argument("--mirror", default=ML.LIVE, help="mirror live folder (default: sharepoint-lists/mirror/live)")
    ap.add_argument("--verbose", action="store_true")
    ap.add_argument("--strict", action="store_true", help="missing new tables / BO columns FAIL (go-live gate)")
    ap.add_argument("--first-build", action="store_true", help="list tables must equal the mirror exactly (no extra Ids)")
    ap.add_argument("--today", type=lambda s: dt.date.fromisoformat(s), help="Eastern 'today' for the Nightly Sync replay")
    opts = ap.parse_args(argv)
    base = Book.from_xlsx(opts.base)
    new = base if os.path.abspath(opts.new) == os.path.abspath(opts.base) else Book.from_xlsx(opts.new)
    mirror = Mirror.from_live(opts.mirror)
    results = run_checks(base, new, mirror, Vocab(), opts)
    print(render_console(results, opts.verbose))
    if opts.report:
        with open(opts.report, "w", encoding="utf-8", newline="\n") as f:
            f.write(render_report(results, base, new, mirror, opts))
        print("report:", opts.report)
    return 1 if any(r.status == "FAIL" for r in results) else 0


if __name__ == "__main__":
    sys.exit(main())
