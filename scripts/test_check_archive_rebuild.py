#!/usr/bin/env python3
"""check_archive_rebuild.py must PASS a good rebuild and FAIL on each unsafe mutation.

    python scripts/test_check_archive_rebuild.py            # synthetic + xlsx round trip + real BASE vs BASE
    python scripts/test_check_archive_rebuild.py --no-real  # skip the real baseline run

Synthetic data is built in memory (a tiny mirror, a 93-column baseline, a 111-column "good" rebuild with
the five list tables); one pass also goes through real .xlsx files with table parts so the zip/openpyxl
readers are exercised. The real run compares the 09-27 23:56 baseline with itself against the live mirror.
"""
import copy
import datetime as dt
import os
import re
import shutil
import sys
import tempfile
import types
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import check_archive_rebuild as C  # noqa: E402

REAL_BASE = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Temp", "claude",
                         "C--Users-solei-OneDrive-Documents-Biplan-claude", "573adf1a-c46f-41dd-9932-bb19548bf9fa",
                         "scratchpad", "archive-work", "Archive active BASE 2356.xlsx")
TODAY = dt.date(2026, 9, 28)
fails = []
VOCAB = C.Vocab()


def opts(**kw):
    o = dict(strict=False, first_build=False, today=TODAY)
    o.update(kw)
    return types.SimpleNamespace(**o)


# ======================================================================= synthetic mirror
OI_FIELDS = sorted(set(C.OI_FIELDS_USED) | {"Id", "Title", "OrderNumberId", "ModelRevisionId", "ItemStatus", "OrderNumber"}
                   | {C.BO_OI_FIELD[c] for c in C.BO_DETAIL})


def oi(i, title, order_id, **kw):
    r = {f: "" for f in OI_FIELDS}
    r.update(Id=str(i), Title=title, OrderNumberId=str(order_id), ModelRevisionId="20", ItemStatus="Active",
             OrderNumber=title.split("-")[0], Qty="2", TemperatureRise="False", Tank="False", CSA="False")
    r.update(kw)
    return r


def mk_mirror(extra_oi_col=None):
    units = [
        oi(1, "30001-1/2", 10, Location="Livraison", ItemStatus="Delivered",
           Planned_x0020_Delivery_x0020_Dat="2026-09-01T04:00:00Z", StepStatus="Terminé",
           StatusDate="2026-09-02T04:00:00Z", CoilingStatus="In Progress", TemperatureRise="True",
           PriceCAD="1234.5", EstimatedDeliveryDate="2026-09-01T04:00:00Z", BO="", BO1PartNumber=""),
        oi(2, "30001-2/2", 10, Location="Bobinage", CoilingDate="2026-08-10T04:00:00Z", CoilingStatus="Completed",
           BO1PartNumber="OI-PART", BO1Date="2026-08-20T04:00:00Z"),
        oi(3, "22140-1/1", 11, Location="Finition", Planned_x0020_Delivery_x0020_Dat="2026-10-05T00:00:00Z"),
    ]
    if extra_oi_col:
        for u in units:
            u[extra_oi_col] = "new"
    orders = [
        dict(Id="10", Title="30001", ClientId="50", Client="ACME", PO="PO-1", Order_x0020_Date="2026-01-15T05:00:00Z",
             Price="1000", EngineeringRequired="True", LDs="False", Initial_x0020_Promised_x0020_Dat="2026-06-01T04:00:00Z",
             Province_x002F_State="QC", WET_x002d_WETP="", Indexing="", ClientDateStatus="", SalesNotes=""),
        dict(Id="11", Title="22140", ClientId="50", Client="ACME", PO="PO-2", Order_x0020_Date="2026-02-01T05:00:00Z",
             Price="", EngineeringRequired="", LDs="True", Initial_x0020_Promised_x0020_Dat="2026-10-06T00:00:00Z",
             Province_x002F_State="QC", WET_x002d_WETP="", Indexing="", ClientDateStatus="", SalesNotes=""),
    ]
    mrs = [dict(Id="20", Title="MR-1", ModelName="G30001", kVA_x0020_and_x0020_kV="500", PrimaryVoltage="",
                SecondaryVoltage="", Phases="3", JS_x0020__x0023_="", Description='["PADMOUNT"]', Model_x0020_Type="PAD",
                Family="", Core_x0020_Type="", Oil_x0020_Type="", OilAmount="", Cable="", Form="",
                Copper_x0028_LV_x0029_="", Wire_x0028_HV_x0029_="", Overcoil="31.75")]
    models = [dict(Id="30", Title="M1")]
    clients = [dict(Id="50", Title="ACME", CliLeadTimeWeeks="20")]
    cat = []
    for lst, rows in (("Order Items", units), ("Order", orders), ("Model Revisions", mrs), ("Models", models), ("Clients", clients)):
        for f in rows[0]:
            typ = "text"
            if f.endswith("Date") or f in ("Planned_x0020_Delivery_x0020_Dat", "Planned_x0020_Tanking_x0020_Date",
                                           "Order_x0020_Date", "Initial_x0020_Promised_x0020_Dat", "EstimatedDeliveryDate"):
                typ = "dateTime (dateOnly)"
            if f in ("Client", "OrderNumber"):
                typ = "lookup"
            cat.append(dict(list=lst, internalName=f, type=typ, csvColumn=f, restField=f, idColumn="",
                            syncedFromList="", syncedFromField="", syncedByFlow="", viaLookup=""))
    return C.Mirror({"Order Items": units, "Order": orders, "Model Revisions": mrs, "Models": models,
                     "Clients": clients, "Columns": cat}, label="synthetic")


MIRROR = mk_mirror()


# ======================================================================= synthetic workbooks
def todt(v):
    return dt.datetime(v.year, v.month, v.day) if isinstance(v, dt.date) and not isinstance(v, dt.datetime) else v


def old_style(v):
    """The viewer-era archive: numbers as invariant text, dates as datetimes."""
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return todt(v)


LEGACY = C.PINNED_FRM1012[82:92]


def frm1012_row(unit, mirror, sync, typed=True):
    ctx = C.context(mirror, unit)
    row = []
    for col in C.PINNED_FRM1012:
        ok, v = C.expected(col, ctx, VOCAB)
        if col in LEGACY:
            v = "Exists" if col == "Archived" else None
        if col == C.SYNC_COL:
            v = sync
        row.append(todt(v) if typed else old_style(v))
    return row


def mk_base():
    units = MIRROR.tables["Order Items"]
    rows = [frm1012_row(u, MIRROR, dt.datetime(2026, 9, 20), typed=False) for u in units]
    h = C.PINNED_FRM1012
    rows[0][h.index("Client")] = "ACME OLD NAME"                           # -> newer from SharePoint
    rows[0][h.index("Lead Time")] = "26"                                   # -> intended (D3)
    rows[2][h.index("Delivery Date")] = dt.datetime(2026, 10, 4)           # -> intended (E7, day early)
    deleted = [None] * len(h)
    deleted[h.index("Order")] = "29999-1/1"
    deleted[h.index("Client")] = "GONE INC"
    deleted[h.index("Location")] = "LI"
    deleted[h.index("Delivery Date")] = dt.datetime(2025, 5, 1)
    deleted[h.index("Tank Delivery Date")] = "."                           # legacy text, kept as is
    deleted[h.index("Archived")] = "Exists"
    rows.append(deleted)
    bo_h = ["Order", "Location", "Status", "Tanking Date", "BO"] + C.BO_DETAIL + [C.SYNC_COL]
    bo_rows = []
    for order, vals in (("30001-1/2", {"BO1 Part Numbre": "BO-PART-1", "BO1 Date": dt.datetime(2026, 8, 1), "BO1 OK": True}),
                        ("30001-2/2", {"BO1 Part Numbre": "OLD-PART", "BO1 OK": False}),
                        ("29999-1/1", {"BO": "late part"})):
        r = [None] * len(bo_h)
        r[0] = order
        for k, v in vals.items():
            r[bo_h.index(k)] = v
        bo_rows.append(r)
    f11_h = ["NUMÉRO DE CUVE", "Date encuvage"] + ["F11 c%d" % i for i in range(37)]
    f13_h = ["Order", "Order Date"] + ["F13 c%d" % i for i in range(53)]
    tables = {
        C.FRM1012: C.Table(C.FRM1012, h, [tuple(r) for r in rows], uid=C.TABLE_GUID, xml_id="2", part="xl/tables/table2.xml", sheet="Archive FRM10-12", ref="A1:CO5"),
        C.BO: C.Table(C.BO, bo_h, [tuple(r) for r in bo_rows], uid="{BO}", ref="A1:X4"),
        C.FRM11: C.Table(C.FRM11, f11_h, [("30001-1/2", dt.datetime(2026, 7, 1)) + (None,) * 37, ("old", "re") + (None,) * 37], ref="A1:AM3"),
        C.FRM13: C.Table(C.FRM13, f13_h, [("30001", dt.datetime(2026, 1, 15)) + (None,) * 53], ref="A1:BC2"),
    }
    return C.Book(tables, label="synthetic base")


def mk_new(base, mirror=MIRROR):
    """A correct rebuild: typed values from the mirror, local rows kept, BO detail merged (D2), five tables."""
    b = base.get(C.FRM1012)
    bk, _ = b.keyed("Order")
    bo = base.get(C.BO)
    bok, _ = bo.keyed("Order")
    oi_t = mirror.oi_by_title()
    rows = []
    for k, brow in bk.items():
        u = oi_t.get(k)
        r = frm1012_row(u, mirror, dt.datetime(2026, 9, 28)) if u else list(brow)
        for col in LEGACY:
            r[C.PINNED_FRM1012.index(col)] = brow[b.idx(col)]
        borow = bok.get(k)
        if not u and borow is not None and C.is_blank(r[b.idx("BO")]):
            r[b.idx("BO")] = borow[bo.idx("BO")]
        for col in C.BO_DETAIL:
            f = C.BO_OI_FIELD[col]
            v = None
            if u and not C.is_blank(u.get(f)):
                v = C.sp_date(u[f]) if col.endswith("Date") else u[f]
            elif borow is not None:
                v = borow[bo.idx(col)]
            r.append(todt(v))
        rows.append(tuple(r))
    t = dict(base.tables)
    t[C.FRM1012] = C.Table(C.FRM1012, C.PINNED_FRM1012 + C.BO_DETAIL, rows, uid=C.TABLE_GUID, xml_id="2",
                           part="xl/tables/table2.xml", sheet="Archive FRM10-12", ref="A1:DG5")
    for tname, lst in C.LIST_TABLES.items():
        cols = mirror.columns(lst)
        dcols = mirror.date_only_cols(lst)
        lrows = [tuple(todt(C.sp_date(r[c])) if c in dcols else (r[c] or None) for c in cols) for r in mirror.tables[lst]]
        t[tname] = C.Table(tname, cols, lrows, uid="{%s}" % tname)
    return C.Book(t, label="synthetic new")


def run(base, new, mirror=MIRROR, **kw):
    return C.run_checks(base, new, mirror, VOCAB, opts(**kw))


def by_num(results, n):
    return [r for r in results if r.num == n][0]


# ======================================================================= mutation helpers
def edit_table(book, name, headers=None, rows=None, **attrs):
    t = book.get(name)
    nt = C.Table(name, headers if headers is not None else t.headers, rows if rows is not None else list(t.rows),
                 uid=t.uid, xml_id=t.xml_id, part=t.part, sheet=t.sheet, ref=t.ref)
    for k, v in attrs.items():
        setattr(nt, k, v)
    tables = dict(book.tables)
    tables[name] = nt
    return C.Book(tables, label=book.label + " (mutated)")


def drop_table(book, name):
    tables = dict(book.tables)
    tables.pop(name)
    return C.Book(tables, label=book.label + " (mutated)")


def set_cell(book, name, key_col, key, col, value):
    t = book.get(name)
    ki, ci = t.idx(key_col), t.idx(col)
    rows = [tuple(value if (j == ci and C.norm_key(r[ki]) == key) else x for j, x in enumerate(r)) for r in t.rows]
    return edit_table(book, name, rows=rows)


def drop_col(book, name, col):
    t = book.get(name)
    i = t.idx(col)
    return edit_table(book, name, headers=t.headers[:i] + t.headers[i + 1:], rows=[r[:i] + r[i + 1:] for r in t.rows])


def swap_cols(book, name, a, b):
    t = book.get(name)
    i, j = t.idx(a), t.idx(b)
    perm = list(range(len(t.headers)))
    perm[i], perm[j] = j, i
    return edit_table(book, name, headers=[t.headers[p] for p in perm], rows=[tuple(r[p] for p in perm) for r in t.rows])


def expect(name, results, check_num, text=None):
    r = by_num(results, check_num)
    if r.has_fail(text):
        msg = [m for l, m, _ in r.items if l == "FAIL"][0]
        print("  ok   %-52s -> [%d] %s" % (name, check_num, msg[:80]))
    else:
        fails.append(name)
        print("  FAIL %-52s -> check %d did not fail" % (name, check_num))


def expect_clean(name, results):
    bad = [(r.num, m) for r in results for l, m, _ in r.items if l == "FAIL"]
    if bad:
        fails.append(name)
        print("  FAIL %-52s -> %s" % (name, bad[:3]))
    else:
        print("  ok   %-52s -> no FAIL (%s)" % (name, " ".join("%d:%s" % (r.num, r.status) for r in results)))


# ======================================================================= xlsx round trip
def write_xlsx(book, path, uids):
    import openpyxl
    from openpyxl.worksheet.table import Table as XTable
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, (name, t) in enumerate(sorted(book.tables.items())):
        ws = wb.create_sheet(("S%d " % i + name)[:31])
        ws.append(list(t.headers))
        for r in t.rows:
            ws.append(list(r))
        ref = "A1:%s%d" % (openpyxl.utils.get_column_letter(len(t.headers)), max(2, len(t.rows) + 1))
        ws.add_table(XTable(displayName=name, ref=ref))
    wb.save(path)
    # openpyxl writes no xr:uid; add it the way Excel stores the table GUID.
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r"xl/tables/table\d+\.xml$", item.filename):
                x = data.decode("utf-8")
                name = re.search(r'displayName="([^"]+)"', x).group(1)
                if name in uids:
                    x = re.sub(r"<table\b", '<table xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" xr:uid="%s"' % uids[name], x, count=1)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def main():
    real = "--no-real" not in sys.argv
    base = mk_base()
    good = mk_new(base)

    print("baseline and good rebuild:")
    expect_clean("BASE vs BASE (synthetic)", run(base, base))
    expect_clean("good rebuild (93 + 18, five tables)", run(base, good))
    r2 = by_num(run(base, good), 2)
    summary = [m for l, m, _ in r2.items if "cells:" in m][0]
    for cls in ("cleaner", "newer", "intended"):
        if " %s 0" % cls in summary:
            fails.append("good rebuild exercises class " + cls)
            print("  FAIL good rebuild never produced a %r cell: %s" % (cls, summary))
    print("       cell classes: %s" % summary)
    expect_clean("good rebuild, --strict", run(base, good, strict=True))

    print("each unsafe mutation must fail:")
    expect("dropped column (Client)", run(base, drop_col(good, C.FRM1012, "Client")), 1, "headers")
    expect("reordered columns (Client <-> KVA and KV)", run(base, swap_cols(good, C.FRM1012, "Client", "KVA and KV")), 1, "headers")
    expect("dropped BO detail column", run(base, drop_col(good, C.FRM1012, "BO2 OK")), 1, "headers")
    expect("text date (Order Date '2026-01-15')", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "Order Date", "2026-01-15")), 3, "text")
    expect("text date is also a cell regression", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "Order Date", "2026-01-15")), 2)
    expect("display-name Location ('Livraison')", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "Location", "Livraison")), 6, "Location")
    t = good.get(C.FRM1012)
    expect("duplicate Order", run(base, edit_table(good, C.FRM1012, rows=list(t.rows) + [t.rows[0]])), 1, "duplicate")
    expect("table recreated (GUID changed)", run(base, edit_table(good, C.FRM1012, uid="{00000000-0000-0000-0000-000000000000}")), 1, "GUID")
    expect("deleted SharePoint unit missing from FRM10-12", run(base, edit_table(good, C.FRM1012, rows=[r for r in t.rows if r[0] != "29999-1/1"])), 5, "lost")
    # second-refresh scenario: the baseline already has the list tables, with an item since deleted on SharePoint
    oi_t = good.get("TableArchiveOrderItems")
    ghost = tuple("99" if h == "Id" else ("ghost-1/1" if h == "Title" else None) for h in oi_t.headers)
    base2 = C.Book(dict(base.tables, TableArchiveOrderItems=C.Table("TableArchiveOrderItems", oi_t.headers, list(oi_t.rows) + [ghost])), "base2")
    expect_clean("second refresh keeps the deleted item (WARN only)", run(base2, edit_table(good, "TableArchiveOrderItems", rows=list(oi_t.rows) + [ghost])))
    expect("deleted SharePoint item missing from list table", run(base2, good), 5, "lost")
    expect("first build with an Id not on SharePoint", run(base, edit_table(good, "TableArchiveOrderItems", rows=list(oi_t.rows) + [ghost])), 4, "no longer")
    expect("mirror Id missing from list table", run(base, edit_table(good, "TableArchiveOrderItems", rows=list(oi_t.rows)[1:])), 4, "missing")
    m2 = mk_mirror(extra_oi_col="NewSharePointColumn")
    expect("new mirror column missing from list table", run(base, good, mirror=m2), 4, "lacks")
    di = oi_t.idx("CoilingDate")
    shifted = [tuple((x - dt.timedelta(days=1)) if (j == di and isinstance(x, dt.datetime)) else x for j, x in enumerate(r)) for r in oi_t.rows]
    expect("list-table date-only shifted a day early", run(base, edit_table(good, "TableArchiveOrderItems", rows=shifted)), 4, "date-only")
    expect("list-table duplicate Id", run(base, edit_table(good, "TableArchiveOrderItems", rows=list(oi_t.rows) + [oi_t.rows[0]])), 4, "duplicate")
    expect("regression value (Client on a live unit)", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "Client", "WRONG")), 2, "REGRESSION")
    expect("regression value (deleted unit changed)", run(base, set_cell(good, C.FRM1012, "Order", "29999-1/1", "Client", "NEW NAME")), 2, "REGRESSION")
    expect("D3 column not SharePoint's value", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "Lead Time", 99)), 2, "REGRESSION")
    expect("lost BO value (BO1 Part Numbre blanked)", run(base, set_cell(good, C.FRM1012, "Order", "30001-1/2", "BO1 Part Numbre", None)), 8, "lost")
    expect("lost BO value on a deleted unit (no backfill)", run(base, set_cell(good, C.FRM1012, "Order", "29999-1/1", "BO", None)), 8, "lost")
    expect("FRM11 cell touched", run(base, set_cell(good, C.FRM11, "NUMÉRO DE CUVE", "old", "Date encuvage", "X")), 1, "TableArchiveFRM11")
    expect("FRM13 column renamed", run(base, edit_table(good, C.FRM13, headers=["Order", "Order date"] + good.get(C.FRM13).headers[2:])), 1, "TableArchiveFRM13")
    expect("list table absent under --strict", run(base, drop_table(good, "TableArchiveClients"), strict=True), 4, "TableArchiveClients")
    expect("reader column gone (Price Value)", run(base, drop_col(good, C.FRM1012, "Price Value")), 6, "PriceReg")

    print("xlsx round trip (zipfile + openpyxl read_only):")
    tmp = tempfile.mkdtemp(prefix="archive-check-")
    try:
        pb, pn, px = (os.path.join(tmp, n) for n in ("base.xlsx", "new.xlsx", "new-recreated.xlsx"))
        write_xlsx(base, pb, {C.FRM1012: C.TABLE_GUID})
        write_xlsx(good, pn, {C.FRM1012: C.TABLE_GUID})
        write_xlsx(good, px, {C.FRM1012: "{11111111-2222-3333-4444-555555555555}"})
        B, N, X = C.Book.from_xlsx(pb), C.Book.from_xlsx(pn), C.Book.from_xlsx(px)
        if B.get(C.FRM1012).headers != C.PINNED_FRM1012 or N.get(C.FRM1012).uid != C.TABLE_GUID:
            fails.append("xlsx reader: headers/GUID")
            print("  FAIL xlsx reader headers/GUID")
        if len(N.get(C.FRM1012).rows) != len(good.get(C.FRM1012).rows):
            fails.append("xlsx reader: row count")
            print("  FAIL xlsx reader row count")
        expect_clean("xlsx: good rebuild", run(B, N))
        expect("xlsx: table recreated (GUID changed)", run(B, X), 1, "GUID")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if real and os.path.exists(REAL_BASE):
        print("real baseline vs itself, live mirror:")
        B = C.Book.from_xlsx(REAL_BASE)
        res = C.run_checks(B, B, C.Mirror.from_live(), VOCAB, opts(today=None))
        expect_clean("REAL BASE vs BASE", res)
    elif real:
        print("  skip real baseline (not found: %s)" % REAL_BASE)

    print("\n%s" % ("ALL PASS" if not fails else "%d FAILED: %s" % (len(fails), fails)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
